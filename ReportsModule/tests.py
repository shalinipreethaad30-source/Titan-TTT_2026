from django.test import TestCase

from modelmasterapp.models import ModelMaster

from .selectors import search_plating_stock


class PlatingStockAutocompleteTests(TestCase):
	def test_search_plating_stock_matches_catalogue_values_without_exact_spacing(self):
		ModelMaster.objects.create(
			model_no='M-001',
			ep_bath_type='EP',
			version='V1',
			plating_stk_no='2617SAA02',
		)

		self.assertEqual(search_plating_stock('2617 SAA'), ['2617SAA02'])


"""Report mapping regressions; no production database access required."""
from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase, RequestFactory

from . import selectors as report
from .selectors import JourneyRecords, submission_values
from modelmasterapp.models import TotalStockModel
from Jig_Loading.models import JigCompleted
from Jig_Unloading.models import JigUnloadAfterTable


IN = datetime(2026, 9, 7, 5, 42, tzinfo=timezone.utc)
OUT = IN + timedelta(minutes=7)


def evidence():
    return JourneyRecords([])


def snapshot(stage, done=True):
    values = dict(is_completed=done, is_draft=not done, created_at=OUT,
                  submission_type='FULL_ACCEPT', total_lot_qty=95,
                  accepted_qty=95, rejected_qty=0)
    if stage == 'IQF':
        values['iqf_incoming_qty'] = 95
    return SimpleNamespace(**values)


def stage_cell(stage, state):
    records = evidence()
    if stage == 'Day Planning':
        batch = None if state == 'absent' else SimpleNamespace(
            Moved_to_D_Picker=state == 'done', date_time=IN,
            total_batch_quantity=95, dp_pick_remarks='')
        return report._day_planning_cell(batch)[0]
    if stage in dict(report._EARLY_MODULE_SPECS):
        stock = TotalStockModel(lot_id='L1', total_stock=95,
                                total_IP_accpeted_quantity=95)
        if state != 'absent':
            records.entries[stage]['L1'] = {'in_time': IN, 'lot_qty': 95}
        if state == 'done':
            if stage == 'Input Screening':
                sub = SimpleNamespace(is_submitted=True, Draft_Saved=False,
                    is_full_accept=True, is_full_reject=False,
                    original_lot_qty=95, submitted_at=OUT)
            else:
                sub = snapshot(stage)
            records.submissions[stage]['L1'] = sub
        return report._early_module_cells([stock], records=records)[0][stage]
    if stage in ('Jig Loading', 'IP Inspection'):
        jig = None
        if state != 'absent':
            records.entries['Jig Loading']['L1'] = {'in_time': IN, 'lot_qty': 95}
            is_submit = stage == 'IP Inspection' or state == 'done'
            records.submissions['Jig Loading']['L1'] = SimpleNamespace(
                status_flag='SUBMITTED' if is_submit else 'DRAFT',
                lot_qty=95, loaded_cases_qty=95, updated_at=IN if stage == 'IP Inspection' else OUT)
            jig = JigCompleted(lot_id='L1', draft_status='submitted' if is_submit else 'draft',
                               loaded_cases_qty=95,
                               jig_position='P1' if stage == 'IP Inspection' and state == 'done' else None,
                               IP_loaded_date_time=OUT)
        result = report._jig_loading_cells(jig, records=records, lot_id='L1')
        return result[0 if stage == 'Jig Loading' else 1]
    zone = stage[-2:].lower()
    unload = None if state == 'absent' else JigUnloadAfterTable(
        lot_id='U1', plating_color_id=1, total_case_qty=95, created_at=IN)
    if unload:
        if stage.startswith('Jig Unloading'):
            if state == 'done':
                unload.Un_loaded_date_time = OUT
        elif stage.startswith('Spider'):
            records.entries[stage]['U1'] = {'in_time': IN}
            if state == 'done':
                setattr(unload, f'ss_{zone}_completed', True)
                setattr(unload, f'ss_{zone}_completed_at', OUT)
        else:
            base = stage[:-3]
            records.entries[base]['U1'] = {'in_time': IN, 'lot_qty': 95}
            if state == 'done':
                records.submissions[base]['U1'] = snapshot(base)
    return report._late_module_cells(unload, {1: zone}, records=records)[0][stage]


class ConsolidatedStageTests(SimpleTestCase):
    def test_in_progress_never_borrows_previous_out(self):
        stock = TotalStockModel(lot_id='L1', total_stock=95, current_stage='IQF',
                                total_IP_accpeted_quantity=55)
        cells, statuses, _ = report._early_module_cells([stock], OUT, evidence())
        self.assertEqual(statuses['IQF'], 'In Progress')
        self.assertIn('IN : --', cells['IQF'])
        self.assertIn('Lot Qty : 55', cells['IQF'])

    def test_iqf_submission_overrides_mutable_stock_quantity(self):
        records = evidence()
        records.submissions['IQF']['L1'] = snapshot('IQF')
        records.submissions['IQF']['L1'].iqf_incoming_qty = 55
        records.submissions['IQF']['L1'].accepted_qty = 55
        stock = TotalStockModel(lot_id='L1', total_stock=150)
        cells = report._early_module_cells([stock], records=records)[0]
        self.assertIn('Lot Qty : 55', cells['IQF'])
        self.assertIn('Rejected : 0', cells['IQF'])

    def test_incomplete_submission_suppresses_out_and_result_quantities(self):
        for stage in ('IQF', 'Brass QC', 'Brass Audit', 'Nickel Wiping', 'Nickel Audit'):
            with self.subTest(stage=stage):
                values = submission_values(stage, snapshot(stage, False))
                self.assertEqual(values['status'], 'In Progress')
                self.assertIsNone(values['out_time'])
                self.assertNotIn('accepted_qty', values)

    def test_refresh_order_does_not_select_inherited_child_flags(self):
        records = evidence()
        records.submissions['Brass QC']['PARENT'] = snapshot('Brass QC')
        parent = TotalStockModel(lot_id='PARENT', total_stock=95)
        child = TotalStockModel(lot_id='CHILD', total_stock=5,
                                brass_qc_rejection=True, bq_last_process_date_time=OUT + timedelta(hours=1))
        first = report._early_module_cells([parent, child], records=records)
        second = report._early_module_cells([child, parent], records=records)
        self.assertEqual(first, second)
        self.assertIn('Lot Qty : 95', first[0]['Brass QC'])

    def test_unload_draft_without_after_table_is_in_progress(self):
        records = evidence()
        records.entries['Jig Unloading']['L1'] = {'in_time': IN, 'lot_qty': 95}
        for zone in ('z1', 'z2'):
            cell = report._late_module_cells(None, {1: zone}, records=records,
                lot_id='L1', plating_color_id=1)[0][f'Jig Unloading {zone.upper()}']
            self.assertIn('Status : In Progress', cell)
            self.assertIn('Lot Qty : 95', cell)

    def test_not_reached_styling(self):
        self.assertEqual(report._stage_state('Not Reached'), report.STATE_NOT_REACHED)

    def test_combined_loading_uses_each_lots_saved_allocation(self):
        records = evidence()
        record = SimpleNamespace(lot_id='L1', lot_qty=150, loaded_cases_qty=120,
            created_at=IN, updated_at=OUT, status_flag='SUBMITTED',
            multi_model_allocation=[{'lot_id': 'L1', 'requested_qty': 95, 'allocated_qty': 80},
                                    {'lot_id': 'L2', 'requested_qty': 55, 'allocated_qty': 40}])
        records._loading_record(record, {'L1', 'L2'})
        self.assertEqual(records.entry('Jig Loading', 'L2')['lot_qty'], 55)
        values = submission_values('Jig Loading', records.submission('Jig Loading', 'L2'))
        self.assertEqual(values['accepted_qty'], 40)

    def test_partial_input_screening_uses_saved_child_quantities(self):
        records = evidence()
        records.submissions['Input Screening']['L1'] = SimpleNamespace(
            is_submitted=True, Draft_Saved=False, is_full_accept=False,
            is_full_reject=False, original_lot_qty=95, submitted_at=OUT)
        records.is_quantities['L1'] = {'accepted_qty': 80, 'rejected_qty': 15}
        stock = TotalStockModel(lot_id='L1', total_stock=10)
        cell = report._early_module_cells([stock], records=records)[0]['Input Screening']
        self.assertIn('Accepted : 80', cell)
        self.assertIn('Rejected : 15', cell)

    def test_input_screening_child_does_not_add_false_in_progress_block(self):
        records = evidence()
        parent = TotalStockModel(lot_id='PARENT', total_stock=168)
        child = TotalStockModel(lot_id='CHILD', total_stock=167,
                                last_process_module='Input Screening',
                                next_process_module='Brass QC')
        records.submissions['Input Screening']['PARENT'] = SimpleNamespace(
            is_submitted=True, Draft_Saved=False, is_full_accept=False,
            is_full_reject=False, original_lot_qty=168, submitted_at=OUT,
        )
        records.is_quantities['PARENT'] = {'accepted_qty': 167, 'rejected_qty': 1}
        cells, statuses, _ = report._early_module_cells([parent, child], records=records)
        self.assertEqual(cells['Input Screening'].count('Status :'), 1)
        self.assertEqual(statuses['Input Screening'], 'Partially Accepted')


    def test_dp_transfer_gate_before_any_input_screening_scan(self):
        from modelmasterapp.models import ModelMasterCreation
        batch = ModelMasterCreation(Moved_to_D_Picker=False, date_time=IN,
                                    total_batch_quantity=95)
        stock = TotalStockModel(lot_id='L1', batch_id=batch, total_stock=95,
                                tray_scan_status=True, next_process_module='IP Screening')
        records = evidence()
        records.add_stock_receipts([stock])
        self.assertEqual(report._day_planning_cell(batch)[1], 'In Progress')
        self.assertIsNone(report._early_module_cells([stock], records=records)[1]['Input Screening'])
        batch.Moved_to_D_Picker = True
        records.add_stock_receipts([stock])
        cells, statuses, _ = report._early_module_cells([stock], records=records)
        self.assertEqual(report._day_planning_cell(batch)[1], 'Completed')
        self.assertEqual(statuses['Input Screening'], 'In Progress')
        self.assertIn('Lot Qty : 95', cells['Input Screening'])
        # The transfer flag has no persisted timestamp: do not reuse DP IN.
        self.assertIn('IN : --', cells['Input Screening'])

    def test_brass_audit_receipt_starts_jig_before_first_draft(self):
        stock = TotalStockModel(lot_id='L1', total_stock=95,
                                brass_audit_accptance=True, brass_audit_accepted_qty=95)
        records = evidence()
        records.add_stock_receipts([stock])
        result = report._jig_loading_cells(None, records=records, lot_id='L1')
        self.assertEqual(result[2], 'In Progress')
        self.assertIn('Lot Qty : 95', result[0])
        self.assertIsNone(result[3])

    def test_ip_handoff_starts_unloading_before_first_draft(self):
        jig = JigCompleted(last_process_module='Inprocess Inspection',
                           loaded_cases_qty=95, IP_loaded_date_time=OUT)
        for zone in ('z1', 'z2'):
            result = report._late_module_cells(None, {1: zone}, records=evidence(),
                                               plating_color_id=1, jig_record=jig)
            self.assertEqual(result[1][f'Jig Unloading {zone.upper()}'], 'In Progress')
            self.assertIn('Lot Qty : 95', result[0][f'Jig Unloading {zone.upper()}'])

    def test_nickel_and_spider_receipts_do_not_wait_for_local_scans(self):
        for zone in ('z1', 'z2'):
            unload = JigUnloadAfterTable(lot_id='U1', total_case_qty=95, plating_color_id=1)
            result = report._late_module_cells(unload, {1: zone}, records=evidence())
            self.assertEqual(result[1][f'Nickel Wiping {zone.upper()}'], 'In Progress')
            self.assertIsNone(result[1][f'Nickel Audit {zone.upper()}'])
            unload.nq_qc_accptance = True
            result = report._late_module_cells(unload, {1: zone}, records=evidence())
            self.assertEqual(result[1][f'Nickel Audit {zone.upper()}'], 'In Progress')
            unload.na_qc_accptance = True
            result = report._late_module_cells(unload, {1: zone}, records=evidence())
            self.assertEqual(result[1][f'Spider Spindle {zone.upper()}'], 'In Progress')
            other = 'Z2' if zone == 'z1' else 'Z1'
            self.assertIsNone(result[1][f'Spider Spindle {other}'])

    def test_zone_two_nickel_wiping_submission_completes_jig_unloading(self):
        unload = JigUnloadAfterTable(lot_id='U2', total_case_qty=78, plating_color_id=2)
        records = evidence()
        wiping = snapshot('Nickel Wiping')
        records.submissions['Nickel Wiping']['U2'] = wiping
        records.submission_history['Nickel Wiping']['U2'].append(wiping)
        cells, statuses, _ = report._late_module_cells(unload, {2: 'z2'}, records=records)
        self.assertEqual(statuses['Jig Unloading Z2'], 'Completed')
        self.assertIn('OUT: ' + report._fmt(OUT), cells['Jig Unloading Z2'])

    def test_new_audit_cycle_ignores_old_rejection(self):
        unload = JigUnloadAfterTable(lot_id='U1', total_case_qty=95, plating_color_id=1,
            na_qc_rejection=True, na_last_process_date_time=IN,
            nq_qc_accptance=True, nq_last_process_date_time=OUT)
        records = evidence()
        old = snapshot('Nickel Audit')
        old.created_at = IN
        old.submission_type = 'FULL_REJECT'
        records.submissions['Nickel Audit']['U1'] = old
        records.entries['Nickel Audit']['U1'] = {'in_time': IN, 'lot_qty': 95}
        cells, statuses, _ = report._late_module_cells(unload, {1: 'z1'}, records=records)
        self.assertEqual(statuses['Nickel Audit Z1'], 'In Progress')
        self.assertIn('OUT: --', cells['Nickel Audit Z1'])

    def test_new_planning_batch_is_visible_before_stock_creation(self):
        from modelmasterapp.models import ModelMasterCreation, Plating_Color
        batch = ModelMasterCreation(date_time=IN, total_batch_quantity=95,
                                    plating_stk_no='NEW-PLAN')
        with patch.object(TotalStockModel, 'objects'), patch.object(
                JigCompleted, 'objects'), patch.object(JigUnloadAfterTable, 'objects'), patch.object(
                Plating_Color, 'objects'), patch.object(ModelMasterCreation, 'objects') as manager:
            manager.filter.return_value.exclude.return_value.order_by.return_value.__iter__.return_value = [batch]
            rows = report.get_consolidated_report_rows()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['module_states']['Day Planning'], report.STATE_CURRENT)
        self.assertEqual(rows[0]['module_states']['Input Screening'], report.STATE_NOT_REACHED)
        self.assertIn('Lot Qty : 95', rows[0]['modules']['Day Planning'])


    def test_input_screening_completion_starts_brass_qc_only(self):
        stock = TotalStockModel(lot_id='L1', total_stock=95,
                                next_process_module='Brass QC')
        records = evidence()
        records.entries['Input Screening']['L1'] = {'in_time': IN, 'lot_qty': 95}
        records.submissions['Input Screening']['L1'] = SimpleNamespace(
            is_submitted=True, Draft_Saved=False, is_full_accept=True,
            is_full_reject=False, original_lot_qty=95, submitted_at=OUT)
        records.entries['Brass QC']['L1'] = {'in_time': OUT, 'lot_qty': 95}
        cells, statuses, _ = report._early_module_cells([stock], records=records)
        self.assertEqual(statuses['Input Screening'], 'Accepted')
        self.assertEqual(statuses['Brass QC'], 'In Progress')
        self.assertIn('IN : ' + report._fmt(OUT), cells['Brass QC'])
        self.assertIn('OUT: --', cells['Brass QC'])
        self.assertIsNone(statuses['IQF'])
        self.assertIsNone(statuses['Brass Audit'])


    def test_dp_transaction_time_is_shared_with_is_receipt(self):
        from modelmasterapp.models import ModelMasterCreation
        batch = ModelMasterCreation(pk=1, date_time=IN, total_batch_quantity=90,
                                    Moved_to_D_Picker=True)
        stock = TotalStockModel(lot_id='L1', batch_id=batch, total_stock=90,
                                tray_scan_status=True)
        records = evidence()
        records.dp_transfers['L1'] = {'transfer_time': OUT, 'lot_qty': 90}
        records.add_stock_receipts([stock])
        dp = report._day_planning_cell(batch, OUT)[0]
        cells = report._early_module_cells([stock], records=records)[0]
        self.assertIn('OUT: ' + report._fmt(OUT), dp)
        self.assertIn('IN : ' + report._fmt(OUT), cells['Input Screening'])
        self.assertIn('OUT: --', cells['Input Screening'])
        batch.Moved_to_D_Picker = False
        self.assertIn('OUT: --', report._day_planning_cell(batch, OUT)[0])

    def test_confirmed_early_routes_use_transfer_timestamps(self):
        cases = [
            ('Input Screening', 'Brass QC', 'last_process_date_time'),
            ('Brass QC', 'IQF', 'bq_last_process_date_time'),
            ('Brass QC', 'Brass Audit', 'bq_last_process_date_time'),
            ('IQF', 'Brass QC', 'iqf_last_process_date_time'),
            ('Brass Audit', 'Jig Loading', 'brass_audit_last_process_date_time'),
        ]
        for source, destination, field in cases:
            with self.subTest(source=source, destination=destination):
                records = evidence()
                stock = TotalStockModel(lot_id='L1', total_stock=95,
                    last_process_module=source, next_process_module=destination,
                    total_IP_accpeted_quantity=95, **{field: OUT})
                records.add_stock_receipts([stock])
                self.assertEqual(records.entry(destination, 'L1')['in_time'], OUT)
                # A later local draft must not shift the original receipt time.
                records.entries[destination]['L1']['in_time'] = OUT + timedelta(hours=1)
                records.add_stock_receipts([stock])
                self.assertEqual(records.entry(destination, 'L1')['in_time'], OUT)

    def test_brass_audit_receipt_uses_brass_qc_accepted_quantity(self):
        records = evidence()
        stock = TotalStockModel(
            lot_id='L1', total_stock=90, brass_qc_accepted_qty=86,
            last_process_module='Brass QC', next_process_module='Brass Audit',
            bq_last_process_date_time=OUT,
        )
        records.add_stock_receipts([stock])
        self.assertEqual(records.entry('Brass Audit', 'L1')['lot_qty'], 86)

    def test_late_transfer_timestamps_before_local_drafts(self):
        for zone in ('z1', 'z2'):
            jig = JigCompleted(last_process_module='Inprocess Inspection',
                               IP_loaded_date_time=OUT, loaded_cases_qty=95)
            cells = report._late_module_cells(None, {1: zone}, records=evidence(),
                plating_color_id=1, jig_record=jig)[0]
            self.assertIn('IN : ' + report._fmt(OUT), cells[f'Jig Unloading {zone.upper()}'])
            unload = JigUnloadAfterTable(lot_id='U1', plating_color_id=1,
                total_case_qty=95, created_at=IN, Un_loaded_date_time=OUT)
            cells = report._late_module_cells(unload, {1: zone}, records=evidence())[0]
            self.assertIn('IN : ' + report._fmt(OUT), cells[f'Nickel Wiping {zone.upper()}'])
            unload.nq_qc_accptance = True
            unload.nq_last_process_date_time = OUT + timedelta(minutes=1)
            unload.na_qc_accptance = True
            unload.na_last_process_date_time = OUT + timedelta(minutes=2)
            cells = report._late_module_cells(unload, {1: zone}, records=evidence())[0]
            self.assertIn('IN : ' + report._fmt(unload.nq_last_process_date_time),
                          cells[f'Nickel Audit {zone.upper()}'])
            self.assertIn('IN : ' + report._fmt(unload.na_last_process_date_time),
                          cells[f'Spider Spindle {zone.upper()}'])

    def test_jig_unloading_keeps_ip_inspection_quantity(self):
        unload = JigUnloadAfterTable(
            lot_id='U1', total_case_qty=142, plating_color_id=1,
            created_at=IN, Un_loaded_date_time=OUT,
        )
        jig = JigCompleted(loaded_cases_qty=144)
        cells, statuses, _ = report._late_module_cells(
            unload, {1: 'z1'}, records=evidence(), jig_record=jig)
        self.assertEqual(statuses['Jig Unloading Z1'], 'Completed')
        self.assertIn('Lot Qty : 144', cells['Jig Unloading Z1'])

    def test_preview_download_share_cell_values_and_permissions(self):
        from . import views
        from openpyxl import load_workbook
        modules = {name: stage_cell(name, 'active') for name in report.MODULE_COLUMNS}
        row = {'s_no': 1, 'plating_stk_no': 'TEST', 'lot_qty': 95,
               'modules': modules, 'remarks': ''}
        request = RequestFactory().get('/?plating_stk_no=TEST')
        request.user = SimpleNamespace(is_authenticated=True)
        with patch('adminportal.services.is_admin_user', return_value=True), patch.object(
                report, 'get_consolidated_report_rows', return_value=[row]):
            first = views.consolidated_report_preview(request)
            second = views.consolidated_report_preview(request)
            download = views.consolidated_report_download(request)
        self.assertEqual(first.content, second.content)
        self.assertEqual(json.loads(first.content)['results'][0]['modules'], modules)
        workbook = load_workbook(BytesIO(download.content))
        rows = list(workbook.active.values)
        values = dict(zip(rows[0], rows[1]))
        for stage, cell in modules.items():
            self.assertEqual(values[stage], cell)
        with patch('adminportal.services.is_admin_user', return_value=False):
            self.assertEqual(views.consolidated_report_preview(request).status_code, 403)
            self.assertEqual(views.consolidated_report_download(request).status_code, 403)


def _matrix_test(stage, state):
    def test(self):
        cell = stage_cell(stage, state)
        details = {line['label']: line['value'] for line in report._parse_cell_lines(cell)}
        if state == 'absent':
            self.assertEqual(details['Status'], 'Not Reached')
            self.assertEqual(details['IN'], '--')
            self.assertEqual(details['OUT'], '--')
            self.assertEqual(details['Lot Qty'], '--')
        else:
            self.assertEqual(details['IN'], report._fmt(IN))
            self.assertEqual(details['Lot Qty'], '95')
            if state == 'active':
                self.assertEqual(details['Status'], 'In Progress')
                self.assertEqual(details['OUT'], '--')
            else:
                self.assertIn(details['Status'], ('Completed', 'Accepted'))
                # DP has no stored completion field; do not substitute IN.
                self.assertEqual(details['OUT'], '--' if stage == 'Day Planning' else report._fmt(OUT))
    return test


for _stage in report.MODULE_COLUMNS:
    for _state in ('absent', 'active', 'done'):
        setattr(ConsolidatedStageTests, 'test_' + _stage.lower().replace(' ', '_') + '_' + _state,
                _matrix_test(_stage, _state))


class ConsolidatedRetrievalTests(SimpleTestCase):
    def planning_rows(self, count=25, **filters):
        from modelmasterapp.models import ModelMasterCreation, Plating_Color
        batches = [ModelMasterCreation(pk=i+1, date_time=IN+timedelta(days=i),
                    total_batch_quantity=i+1, plating_stk_no='SAME-STOCK')
                   for i in range(count)]
        with patch.object(TotalStockModel, 'objects'), patch.object(
                JigCompleted, 'objects'), patch.object(JigUnloadAfterTable, 'objects'), patch.object(
                Plating_Color, 'objects'), patch.object(ModelMasterCreation, 'objects') as manager:
            query = manager.filter.return_value.exclude.return_value.order_by.return_value
            query.__iter__.return_value = batches
            query.filter.return_value.__iter__.return_value = batches
            return report.get_consolidated_report_rows(**filters)

    def test_repeated_plating_number_keeps_every_batch_newest_first(self):
        rows = self.planning_rows()
        self.assertEqual(len(rows), 25)
        self.assertEqual([row['lot_qty'] for row in rows], list(range(25, 0, -1)))

    def test_date_filter_keeps_all_matching_batches(self):
        rows = self.planning_rows(date_from=(IN+timedelta(days=3)).date(),
                                  date_to=(IN+timedelta(days=7)).date())
        self.assertEqual([row['lot_qty'] for row in rows], [8, 7, 6, 5, 4])

    def test_preview_paginates_but_export_includes_every_row(self):
        from . import views
        from openpyxl import load_workbook
        rows = self.planning_rows()
        request = RequestFactory().get('/', {'page': 2})
        with patch.object(views, '_consolidated_rows_from_request', return_value=rows):
            preview = json.loads(views.consolidated_report_preview.__wrapped__.__wrapped__(request).content)
            export = views.consolidated_report_download.__wrapped__.__wrapped__(request)
        self.assertEqual(preview['total_records'], 25)
        self.assertEqual(preview['num_pages'], 3)
        self.assertEqual(len(preview['results']), 10)
        self.assertEqual(preview['results'][0]['s_no'], 11)
        workbook = load_workbook(BytesIO(export.content), read_only=True)
        self.assertEqual(workbook.active.max_row-1, 25)
        workbook.close()

class BrassQCTransitionReportTests(SimpleTestCase):
    def test_destination_child_cannot_override_completed_parent(self):
        for decision, accepted, rejected, destination in [
                ('FULL_REJECT', 0, 80, 'IQF'),
                ('FULL_ACCEPT', 80, 0, 'Brass Audit'),
                ('PARTIAL', 50, 30, 'IQF')]:
            with self.subTest(decision=decision):
                records = evidence()
                saved = SimpleNamespace(is_completed=True, submission_type=decision,
                    total_lot_qty=80, accepted_qty=accepted, rejected_qty=rejected,
                    created_at=OUT, transition_lot_id='CHILD')
                records.submissions['Brass QC']['PARENT'] = saved
                records.entries['Brass QC']['PARENT'] = {'in_time': IN, 'lot_qty': 80}
                records.transition_entries['Brass QC']['PARENT'] = {
                    'in_time': IN, 'lot_qty': 80}
                records.entries['Brass QC']['CHILD'] = {'in_time': OUT+timedelta(seconds=1), 'lot_qty': 80}
                parent = TotalStockModel(lot_id='PARENT', total_stock=80)
                child = TotalStockModel(lot_id='CHILD', total_stock=rejected or accepted,
                    current_stage='Brass QC', last_process_module='Brass QC',
                    next_process_module=destination, bq_last_process_date_time=OUT)
                records.rw_quantities['CHILD'] = ((OUT, 1), rejected)
                records.add_stock_receipts([child])
                cells, statuses, _ = report._early_module_cells([parent, child], records=records)
                self.assertEqual(statuses['Brass QC'], {'FULL_REJECT':'Rejected',
                    'FULL_ACCEPT':'Accepted', 'PARTIAL':'Partially Accepted'}[decision])
                self.assertIn('OUT: '+report._fmt(OUT), cells['Brass QC'])
                self.assertIn('Lot Qty : 80', cells['Brass QC'])
                self.assertIn('Rejected : '+str(rejected), cells['Brass QC'])
                if destination == 'IQF':
                    self.assertEqual(statuses['IQF'], 'In Progress')
                    self.assertIn('Lot Qty : '+str(rejected), cells['IQF'])
                child.last_process_module = 'IQF'
                child.current_stage = 'IQF'
                child.iqf_rejection = True
                for destination_after_reject in ('IQF', None):
                    child.next_process_module = destination_after_reject
                    _, after_reject, _ = report._early_module_cells([parent, child], records=records)
                    self.assertEqual(after_reject['Brass QC'], statuses['Brass QC'])
                child.iqf_rejection = False
                child.next_process_module = 'Brass QC'
                returned_cells, statuses, _ = report._early_module_cells([parent, child], records=records)
                self.assertEqual(statuses['Brass QC'], 'In Progress')
                blocks = returned_cells['Brass QC'].split('\n\n')
                self.assertEqual(len(blocks), 2)
                self.assertIn('OUT: '+report._fmt(OUT), blocks[0])
                self.assertIn('Status : In Progress', blocks[1])
                self.assertNotIn('Lot ID', returned_cells['Brass QC'])
                details = report._parse_cell_lines(returned_cells['Brass QC'])
                self.assertEqual([line['block_state'] for line in details if line['label']=='Status'],
                                 ['completed', 'current'])


class ReportRouteApplicabilityTests(SimpleTestCase):
    def test_route_exclusions_preserve_pending_and_actual_stages(self):
        for zone, other in [('z1', 'Z2'), ('z2', 'Z1')]:
            modules = {name: report._module_cell('Not Reached') for name in report.MODULE_COLUMNS}
            statuses = dict.fromkeys(modules)
            statuses['Brass QC'] = 'Accepted'
            report._apply_route_applicability(modules, statuses, zone)
            self.assertEqual(statuses['IQF'], 'Not Applicable')
            for name in ('Jig Unloading', 'Nickel Wiping', 'Nickel Audit'):
                self.assertEqual(statuses[name+' '+other], 'Not Applicable')
                self.assertIsNone(statuses[name+' '+zone.upper()])
            self.assertIsNone(statuses['Brass Audit'])
            self.assertEqual(report._stage_state(statuses['IQF']), report.STATE_NOT_APPLICABLE)
            self.assertIn('Status : Not Applicable', modules['IQF'])

    def test_rejection_and_unknown_route_do_not_exclude_iqf(self):
        for bq_status in (None, 'In Progress', 'Rejected', 'Partially Accepted'):
            modules = {name: report._module_cell('Not Reached') for name in report.MODULE_COLUMNS}
            statuses = dict.fromkeys(modules)
            statuses['Brass QC'] = bq_status
            report._apply_route_applicability(modules, statuses)
            self.assertIsNone(statuses['IQF'])

    def test_real_history_is_never_hidden(self):
        modules = {'IQF': 'recorded IQF', 'Nickel Audit Z2': 'recorded audit'}
        statuses = {'Brass QC': 'Accepted', 'IQF': 'Rejected', 'Nickel Audit Z2': 'Accepted'}
        report._apply_route_applicability(modules, statuses, 'z1')
        self.assertEqual(modules['IQF'], 'recorded IQF')
        self.assertEqual(statuses['Nickel Audit Z2'], 'Accepted')


class BrassQCIncomingQuantityTests(SimpleTestCase):
    def test_input_handoff_uses_accepted_quantity_not_original_stock(self):
        for accepted in (139, 144, 0):
            for saved in (False, True):
                with self.subTest(accepted=accepted, saved=saved):
                    stock = TotalStockModel(lot_id='IS-TO-BQ', total_stock=144,
                        total_IP_accpeted_quantity=accepted, last_process_module='Input Screening',
                        next_process_module='Brass QC', last_process_date_time=OUT)
                    records = evidence()
                    if saved:
                        records.is_quantities[stock.lot_id] = {'accepted_qty': accepted}
                    records.add_stock_receipts([stock])
                    cells, statuses, _ = report._early_module_cells([stock], records=records)
                    self.assertIn('Lot Qty : '+str(accepted), cells['Brass QC'])
                    self.assertEqual(statuses['Brass QC'], 'In Progress')
                    self.assertIn('OUT: --', cells['Brass QC'])


class ConsolidatedSplitBatchTests(SimpleTestCase):
    def test_split_children_share_row_but_separate_batches_do_not(self):
        from modelmasterapp.models import ModelMasterCreation, Plating_Color
        batch = ModelMasterCreation(pk=901, plating_stk_no='REPEATED',
            total_batch_quantity=144, date_time=IN)
        other = ModelMasterCreation(pk=902, plating_stk_no='REPEATED',
            total_batch_quantity=90, date_time=IN)
        stocks = [TotalStockModel(pk=1, lot_id='ACCEPT', batch_id=batch, total_stock=137, created_at=OUT),
                  TotalStockModel(pk=2, lot_id='REJECT', batch_id=batch, total_stock=2, created_at=OUT),
                  TotalStockModel(pk=3, lot_id='OTHER', batch_id=other, total_stock=90, created_at=IN)]
        records = evidence()
        with patch.object(TotalStockModel, 'objects') as manager, patch.object(
                JigCompleted, 'objects'), patch.object(JigUnloadAfterTable, 'objects'), patch.object(
                Plating_Color, 'objects'), patch.object(ModelMasterCreation, 'objects'), patch.object(
                report, 'JourneyRecords', return_value=records):
            manager.filter.return_value.select_related.return_value.order_by.return_value.__iter__.return_value = stocks
            manager.filter.return_value.exclude.return_value.select_related.return_value.order_by.return_value.__iter__.return_value = stocks
            rows = report.get_consolidated_report_rows()
        self.assertEqual(len(rows), 2)
        self.assertEqual([r['lot_qty'] for r in rows], [144, 90])



class BrassAuditBranchDisplayTests(SimpleTestCase):
    def test_independent_audit_lots_keep_quantities_statuses_and_times(self):
        records = evidence()
        stocks = [TotalStockModel(lot_id='AUDIT-137', total_stock=137),
                  TotalStockModel(lot_id='AUDIT-2', total_stock=2),
                  TotalStockModel(lot_id='RETURN-QC', total_stock=2, current_stage='Brass Audit',
                      last_process_module='Brass Audit', next_process_module='Brass QC')]
        records.entries['Brass Audit']['AUDIT-137'] = {'in_time': IN, 'lot_qty': 137}
        records.entries['Brass Audit']['AUDIT-2'] = {'in_time': OUT, 'lot_qty': 2}
        records.submissions['Brass Audit']['AUDIT-2'] = SimpleNamespace(
            is_completed=True, submission_type='FULL_REJECT', total_lot_qty=2,
            accepted_qty=0, rejected_qty=2, created_at=OUT+timedelta(minutes=1))
        cells, statuses, _ = report._early_module_cells(stocks, records=records)
        cell = cells['Brass Audit']
        self.assertNotIn('Lot ID', cell)
        self.assertIn('Lot Qty : 137', cell)
        self.assertLess(cell.index('Status : Rejected'), cell.index('Status : In Progress'))
        self.assertIn('Lot Qty : 2', cell)
        self.assertIn('Status : In Progress', cell)
        self.assertIn('Status : Rejected', cell)
        self.assertNotIn('Lot Qty : 139', cell)
        self.assertNotIn('RETURN-QC', cell)
        self.assertEqual(statuses['Brass Audit'], 'In Progress')
        lines = report._parse_cell_lines(cell)
        self.assertEqual([r['block_state'] for r in lines if r['label']=='Status'], ['completed', 'current'])
        # Repeated stock references must not duplicate a lot's block.
        repeated = report._early_module_cells(stocks+[stocks[0]], records=records)[0]
        self.assertEqual(repeated['Brass Audit'], cell)


class DownstreamBranchBlockTests(SimpleTestCase):
    def test_two_jig_receipts_keep_distinct_quantities(self):
        cells = [report._module_cell('In Progress', in_time=IN, lot_qty=135),
                 report._module_cell('In Progress', in_time=OUT, lot_qty=2)]
        row = {'modules': {}, 'module_details': {}, 'module_states': {},
               '_stage_branches': {'Jig Loading': {
                   ('stock', 1): (cells[0], report.STATE_CURRENT),
                   ('stock', 2): (cells[1], report.STATE_CURRENT)}}}
        report._render_stage_branches(row)
        self.assertEqual(row['modules']['Jig Loading'], '\n\n'.join(cells))
        self.assertNotIn('_stage_branches', row)
        self.assertNotIn('Lot ID', row['modules']['Jig Loading'])
        quantities = [v['value'] for v in row['module_details']['Jig Loading'] if v['label']=='Lot Qty']
        self.assertEqual(quantities, ['135','2'])

    def test_final_branch_precedes_pending_branch(self):
        for stage in ['Jig Loading', 'IP Inspection', 'Jig Unloading Z1', 'Nickel Audit Z2']:
            row = {'modules': {}, 'module_details': {}, 'module_states': {},
                   '_stage_branches': {stage: {
                       1: (report._module_cell('In Progress', lot_qty=135), report.STATE_CURRENT),
                       2: (report._module_cell('Completed', lot_qty=2), report.STATE_COMPLETED)}}}
            report._render_stage_branches(row)
            statuses = [v for v in row['module_details'][stage] if v['label']=='Status']
            self.assertEqual([v['value'] for v in statuses], ['Completed','In Progress'])
            self.assertEqual([v['block_state'] for v in statuses], ['completed','current'])


class EarlyModuleBranchBlockTests(SimpleTestCase):
    def test_input_screening_and_iqf_keep_each_distinct_lot_transaction(self):
        cases = (
            ('Input Screening', 'accepted_Ip_stock', 'total_IP_accpeted_quantity'),
            ('IQF', 'iqf_acceptance', 'iqf_accepted_qty'),
        )
        for stage, accepted_flag, accepted_quantity in cases:
            with self.subTest(stage=stage):
                records = evidence()
                first = TotalStockModel(lot_id=stage + '-A', total_stock=135)
                second = TotalStockModel(lot_id=stage + '-B', total_stock=2)
                setattr(first, accepted_flag, True)
                setattr(second, accepted_flag, True)
                setattr(first, accepted_quantity, 135)
                setattr(second, accepted_quantity, 2)
                records.entries[stage][first.lot_id] = {'in_time': IN, 'lot_qty': 135}
                records.entries[stage][second.lot_id] = {'in_time': OUT, 'lot_qty': 2}
                cells, statuses, _ = report._early_module_cells([first, second], records=records)
                self.assertEqual(cells[stage].count('Status :'), 2)
                self.assertIn('Lot Qty : 135', cells[stage])
                self.assertIn('Lot Qty : 2', cells[stage])
                self.assertNotIn('Lot ID', cells[stage])
                self.assertEqual(statuses[stage], 'Accepted')

class TransitionEntryTimestampTests(SimpleTestCase):
    def test_parent_submission_uses_saved_child_tray_entry_time(self):
        records = evidence()
        parent = TotalStockModel(lot_id='PARENT', total_stock=139)
        saved = snapshot('Brass QC')
        saved.transition_accept_lot_id = 'CHILD'
        records.submissions['Brass QC']['PARENT'] = saved
        records.transition_entries['Brass QC']['PARENT'] = {
            'in_time': IN, 'lot_qty': 139}
        cells, statuses, _ = report._early_module_cells([parent], records=records)
        self.assertEqual(statuses['Brass QC'], 'Accepted')
        self.assertIn('IN : ' + report._fmt(IN), cells['Brass QC'])
        self.assertIn('OUT: ' + report._fmt(OUT), cells['Brass QC'])


class BrassQCEntryTimestampTests(SimpleTestCase):
    def test_brass_audit_submission_uses_brass_qc_handoff_when_entry_is_missing(self):
        records = evidence()
        stock = TotalStockModel(lot_id='L1', total_stock=97)
        qc = snapshot('Brass QC')
        qc.created_at = OUT
        audit = snapshot('Brass Audit')
        audit.submission_type = 'FULL_REJECT'
        audit.created_at = OUT + timedelta(minutes=1)
        records.submissions['Brass QC']['L1'] = qc
        records.submissions['Brass Audit']['L1'] = audit
        cells, statuses, _ = report._early_module_cells([stock], records=records)
        self.assertEqual(statuses['Brass Audit'], 'Rejected')
        self.assertIn('IN : ' + report._fmt(OUT), cells['Brass Audit'])

    def test_completed_qc_without_tray_entry_uses_input_screening_handoff(self):
        records = evidence()
        parent = TotalStockModel(lot_id='PARENT', total_stock=144)
        records.submissions['Input Screening']['PARENT'] = SimpleNamespace(
            is_submitted=True, Draft_Saved=False, is_full_accept=True,
            is_full_reject=False, original_lot_qty=144, submitted_at=OUT,
        )
        records.submissions['Brass QC']['PARENT'] = snapshot('Brass QC')
        cells, statuses, _ = report._early_module_cells([parent], records=records)
        self.assertEqual(statuses['Brass QC'], 'Accepted')
        self.assertIn('IN : ' + report._fmt(OUT), cells['Brass QC'])

    def test_completed_parent_uses_saved_transition_child_tray_entry(self):
        records = evidence()
        parent = TotalStockModel(lot_id='PARENT', total_stock=139)
        accepted_child = TotalStockModel(lot_id='ACCEPTED-CHILD', total_stock=137)
        rejected_child = TotalStockModel(lot_id='REJECTED-CHILD', total_stock=2)
        records.entries['Brass QC']['ACCEPTED-CHILD'] = {
            'in_time': IN, 'lot_qty': 137,
        }
        records.entries['Brass QC']['REJECTED-CHILD'] = {
            'in_time': OUT, 'lot_qty': 2,
        }
        records.submissions['Brass QC']['PARENT'] = SimpleNamespace(
            is_completed=True, submission_type='PARTIAL',
            total_lot_qty=139, accepted_qty=137, rejected_qty=2,
            created_at=OUT, transition_lot_id=None,
            transition_accept_lot_id='ACCEPTED-CHILD',
            transition_reject_lot_id='REJECTED-CHILD',
        )
        cells, statuses, _ = report._early_module_cells(
            [parent, accepted_child, rejected_child], records=records)
        self.assertIn('IN : ' + report._fmt(IN), cells['Brass QC'])
        self.assertIn('OUT: ' + report._fmt(OUT), cells['Brass QC'])
        self.assertEqual(statuses['Brass QC'], 'Partially Accepted')

class NickelTransactionBlockTests(SimpleTestCase):
    def test_audit_rejection_uses_wiping_handoff_when_child_has_no_timestamp(self):
        child = JigUnloadAfterTable(
            lot_id='CHILD', total_case_qty=142, plating_color_id=1,
            nq_qc_accptance=True, na_qc_rejection=True,
            na_last_process_date_time=OUT,
        )
        records = evidence()
        records.nq_partial_accept_parent['CHILD'] = 'PARENT'
        records.nq_partial_accept_parent_unloads['CHILD'] = JigUnloadAfterTable(
            lot_id='PARENT', total_case_qty=144, plating_color_id=1,
            created_at=IN, Un_loaded_date_time=IN,
        )
        wiping = snapshot('Nickel Wiping')
        wiping.created_at = IN
        audit = snapshot('Nickel Audit')
        audit.submission_type = 'FULL_REJECT'
        audit.total_lot_qty = 142
        audit.accepted_qty = 0
        audit.rejected_qty = 142
        audit.created_at = OUT
        records.submissions['Nickel Wiping']['PARENT'] = wiping
        records.submission_history['Nickel Wiping']['PARENT'].append(wiping)
        records.submissions['Nickel Audit']['CHILD'] = audit
        records.submission_history['Nickel Audit']['CHILD'].append(audit)

        cells, statuses, _ = report._late_module_cells(child, {1: 'z1'}, records=records)

        self.assertEqual(statuses['Nickel Audit Z1'], 'Rejected')
        self.assertIn('IN : ' + report._fmt(IN), cells['Nickel Audit Z1'])
        self.assertIn('OUT: ' + report._fmt(OUT), cells['Nickel Audit Z1'])

    def test_wiping_partial_accept_child_uses_parent_transaction(self):
        child = JigUnloadAfterTable(
            lot_id='CHILD', total_case_qty=142, plating_color_id=1,
            nq_qc_accptance=True, nq_qc_accepted_qty=142,
            nq_last_process_date_time=OUT,
        )
        records = evidence()
        records.nq_partial_accept_parent['CHILD'] = 'PARENT'
        records.nq_partial_accept_parent_unloads['CHILD'] = JigUnloadAfterTable(
            lot_id='PARENT', total_case_qty=144, plating_color_id=1,
            created_at=IN, Un_loaded_date_time=OUT,
        )
        partial = snapshot('Nickel Wiping')
        partial.submission_type = 'PARTIAL'
        partial.total_lot_qty = 144
        partial.accepted_qty = 142
        partial.rejected_qty = 2
        records.submissions['Nickel Wiping']['PARENT'] = partial
        records.submission_history['Nickel Wiping']['PARENT'].append(partial)

        cells, statuses, _ = report._late_module_cells(child, {1: 'z1'}, records=records)

        self.assertEqual(statuses['Jig Unloading Z1'], 'Completed')
        self.assertIn('Lot Qty : 144', cells['Jig Unloading Z1'])
        self.assertEqual(statuses['Nickel Wiping Z1'], 'Partially Accepted')
        self.assertIn('Lot Qty : 144', cells['Nickel Wiping Z1'])
        self.assertIn('Accepted : 142', cells['Nickel Wiping Z1'])
        self.assertIn('Rejected : 2', cells['Nickel Wiping Z1'])

    def test_audit_partial_accept_child_is_not_a_new_jig_unloading_pass(self):
        child = JigUnloadAfterTable(
            lot_id='CHILD', total_case_qty=100, plating_color_id=1,
            nq_qc_accptance=True, na_qc_accptance=True,
            na_last_process_date_time=OUT,
        )
        records = evidence()
        records.na_partial_accept_parent['CHILD'] = 'PARENT'
        records.na_partial_accept_parent_unloads['CHILD'] = JigUnloadAfterTable(
            lot_id='PARENT', total_case_qty=144, plating_color_id=1,
            created_at=IN, Un_loaded_date_time=OUT,
        )
        partial = snapshot('Nickel Audit')
        partial.submission_type = 'PARTIAL'
        partial.total_lot_qty = 144
        partial.accepted_qty = 100
        partial.rejected_qty = 44
        records.submissions['Nickel Audit']['PARENT'] = partial
        records.submission_history['Nickel Audit']['PARENT'].append(partial)
        cells, statuses, _ = report._late_module_cells(child, {1: 'z1'}, records=records)
        self.assertEqual(statuses['Jig Unloading Z1'], 'Completed')
        self.assertIn('Lot Qty : 144', cells['Jig Unloading Z1'])
        self.assertEqual(statuses['Nickel Audit Z1'], 'Partially Accepted')
        self.assertIn('Accepted : 100', cells['Nickel Audit Z1'])
        self.assertIn('Rejected : 44', cells['Nickel Audit Z1'])
        self.assertEqual(statuses['Spider Spindle Z1'], 'In Progress')
        self.assertIn('Lot Qty : 100', cells['Spider Spindle Z1'])

    def test_later_audit_acceptance_does_not_reopen_nickel_wiping(self):
        unload = JigUnloadAfterTable(
            lot_id='U1', total_case_qty=144, plating_color_id=1,
            na_qc_rejection=True, na_qc_accptance=True,
            nq_last_process_date_time=OUT,
            na_last_process_date_time=OUT + timedelta(minutes=3),
        )
        records = evidence()
        first_wiping = snapshot('Nickel Wiping')
        first_wiping.created_at = IN
        second_wiping = snapshot('Nickel Wiping')
        second_wiping.created_at = OUT
        rejected_audit = snapshot('Nickel Audit')
        rejected_audit.submission_type = 'FULL_REJECT'
        rejected_audit.created_at = OUT + timedelta(minutes=1)
        accepted_audit = snapshot('Nickel Audit')
        accepted_audit.created_at = OUT + timedelta(minutes=3)
        records.submissions['Nickel Wiping']['U1'] = second_wiping
        records.submission_history['Nickel Wiping']['U1'].extend(
            [first_wiping, second_wiping])
        records.submissions['Nickel Audit']['U1'] = accepted_audit
        records.submission_history['Nickel Audit']['U1'].extend(
            [rejected_audit, accepted_audit])
        cells, statuses, _ = report._late_module_cells(unload, {1: 'z1'}, records=records)
        self.assertEqual(cells['Nickel Wiping Z1'].count('Status :'), 2)
        self.assertNotIn('Status : In Progress', cells['Nickel Wiping Z1'])
        self.assertEqual(statuses['Nickel Audit Z1'], 'Partially Accepted')

    def test_wiping_accept_then_audit_return_keeps_both_blocks(self):
        accepted = dict(status='Accepted', in_time=IN, out_time=OUT,
                        lot_qty=144, accepted_qty=144, rejected_qty=0)
        returned = dict(status='In Progress', in_time=OUT + timedelta(minutes=1),
                        out_time=None, lot_qty=144)
        text, status, _ = report._module_transaction_blocks([accepted], returned)
        self.assertEqual(text.count('Status :'), 2)
        self.assertLess(text.index('Status : Accepted'), text.index('Status : In Progress'))
        self.assertIn('Lot Qty : 144', text)
        self.assertEqual(status, 'In Progress')
        states = [line['block_state'] for line in report._parse_cell_lines(text)
                  if line['label'] == 'Status']
        self.assertEqual(states, ['completed', 'current'])

    def test_nickel_audit_rejection_is_retained_as_final_block(self):
        rejected = dict(status='Rejected', in_time=IN, out_time=OUT,
                        lot_qty=144, accepted_qty=0, rejected_qty=144)
        text, status, _ = report._module_transaction_blocks([rejected], rejected)
        self.assertEqual(text.count('Status :'), 1)
        self.assertIn('Status : Rejected', text)
        self.assertEqual(status, 'Rejected')
