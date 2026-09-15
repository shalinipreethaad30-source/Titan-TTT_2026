from django.shortcuts import render
from django.views.generic import TemplateView
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils.decorators import method_decorator
from adminportal.decorators import require_admin
from modelmasterapp.models import *
from Recovery_DP.models import *
from DayPlanning.models import DPTrayId_History
from InputScreening.models import IPTrayId, IP_Accepted_TrayScan, IP_Rejected_TrayScan, IP_Accepted_TrayID_Store
from Brass_QC.models import BrassTrayId, Brass_Qc_Accepted_TrayScan, Brass_QC_Rejected_TrayScan, Brass_Qc_Accepted_TrayID_Store, Brass_QC_Rejection_ReasonStore, Brass_QC_Draft_Store
from IQF.models import IQFTrayId, IQF_Accepted_TrayScan, IQF_Rejected_TrayScan, IQF_Accepted_TrayID_Store
from BrassAudit.models import BrassAuditTrayId, Brass_Audit_Accepted_TrayScan, Brass_Audit_Rejected_TrayScan, Brass_Audit_Accepted_TrayID_Store
from django.core.paginator import Paginator
from django.db.models import OuterRef, Subquery, Sum, IntegerField, Count
from django.utils import timezone
import logging
import pytz
from datetime import timedelta
import pandas as pd
from io import BytesIO
from datetime import datetime

logger = logging.getLogger(__name__)

def _autofit_excel_columns(writer, max_width=60):
    """
    Auto-fit every column in every sheet of the workbook so no value is
    cropped. Multi-line cells (e.g. Current/Next Stage) get wrap_text and
    are sized to their longest line. Must be called inside the
    pd.ExcelWriter context, after all sheets are written.
    """
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Alignment

    for ws in writer.book.worksheets:
        for idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            has_multiline = False
            for cell in column_cells:
                if cell.value is None:
                    continue
                text = str(cell.value)
                if '\n' in text:
                    has_multiline = True
                    line_len = max(len(line) for line in text.split('\n'))
                else:
                    line_len = len(text)
                max_len = max(max_len, line_len)
            ws.column_dimensions[get_column_letter(idx)].width = min(
                max_len + 2, max_width
            )
            if has_multiline:
                for cell in column_cells:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')


def _workbook_has_data(writer):
    """True when at least one sheet has a data row (row 1 is the header)."""
    return any(ws.max_row > 1 for ws in writer.book.worksheets)


def convert_datetimes(data):
    for item in data:
        for key, value in item.items():
            if isinstance(value, datetime) and value.tzinfo is not None:
                item[key] = value.replace(tzinfo=None)
    return data

@method_decorator(login_required(login_url='login'), name='dispatch')
@method_decorator(require_admin, name='dispatch')
class ReportsView(TemplateView):
    template_name = "reports.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # List of available modules for reports
        context['modules'] = [
            {'value': 'day-planning', 'label': 'Day Planning'},
            {'value': 'input-screening', 'label': 'Input Screening'},
            {'value': 'brass-qc', 'label': 'Brass QC'},
            {'value': 'iqf', 'label': 'IQF'},
            {'value': 'brass-audit', 'label': 'Brass Audit'},
            {'value': 'recovery-day-planning', 'label': 'Recovery Day Planning'},
            {'value': 'recovery-input-screening', 'label': 'Recovery Input Screening'},
            {'value': 'recovery-brass-qc', 'label': 'Recovery Brass QC'},
            {'value': 'recovery-iqf', 'label': 'Recovery IQF'},
            {'value': 'recovery-brass-audit', 'label': 'Recovery Brass Audit'},
            {'value': 'jig-loading', 'label': 'Jig Loading'},
            {'value': 'inprocess-inspection', 'label': 'Inprocess Inspection'},
            {'value': 'jig-unloading', 'label': 'Jig Unloading'},
            {'value': 'nickel-inspection', 'label': 'Nickel Inspection'},
            {'value': 'nickel-audit', 'label': 'Nickel Audit'},
            {'value': 'spider-spindle', 'label': 'Spider Spindle'}
        ]
        return context
    
# Function for "Reports Module" to download Excel report based on selected module
@login_required(login_url='login')
@require_admin
def download_report(request):
    module = request.GET.get('module')
    if not module:
        return HttpResponse("Module not specified", status=400)

    # Create Excel file with multiple sheets
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        if module == 'day-planning':
            # Custom format for Day Planning report
            from modelmasterapp.models import ModelMasterCreation
            batches = ModelMasterCreation.objects.filter(total_batch_quantity__gt=0, Moved_to_D_Picker=False).select_related('location', 'version').annotate(
                next_process_module=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('next_process_module')[:1])
            )
            
            report_data = []
            for idx, batch in enumerate(batches, start=1):
                # Determine Lot Status - matching Day Planning pick table values
                if batch.hold_lot:
                    lot_status = "On Hold"
                elif batch.next_process_module:
                    lot_status = "Released"
                else:
                    lot_status = "Yet to Start"
                
                # Tray Cate-Capacity
                tray_cate_capacity = f"{batch.tray_type} - {batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ""
                
                # Combine holding and releasing reasons
                remarks_holding = f"Holding Reason: {batch.holding_reason}" if batch.holding_reason else ""
                remarks_releasing = f"Releasing Reason: {batch.release_reason}" if batch.release_reason else ""
                combined_remarks = "\n".join(filter(None, [remarks_holding, remarks_releasing])) or ""
                
                #Day Planning Excel Columns
                row = {
                    'S.No': idx,
                    'Date & Time': batch.date_time.replace(tzinfo=None) if batch.date_time else None,
                    'Plating Stock No': batch.plating_stk_no or '',
                    'Plating color': batch.plating_color or '',
                    'Lot Status': lot_status,
                    'Remarks (for holding row)': combined_remarks,
                    'Category': batch.category or '',
                    'Tray Cate-Capacity': tray_cate_capacity,
                    #'No of Trays': batch.no_of_trays or 0,
                    'Input Qty': batch.total_batch_quantity or 0,
                    'Current Stage': batch.next_process_module or 'Day Planning',
                    'Remarks (chat)': batch.dp_pick_remarks or ''
                }
                report_data.append(row)
            
            df = pd.DataFrame(report_data)
            df.to_excel(writer, sheet_name='Day Planning Report', index=False)
            
            # Completed Table
            from django.utils import timezone
            tz = pytz.timezone("Asia/Kolkata")
            now_local = timezone.now().astimezone(tz)
            today = now_local.date()
            yesterday = today - timedelta(days=1)
            from_date = yesterday
            to_date = today
            from_datetime = timezone.make_aware(datetime.combine(from_date, datetime.min.time()))
            to_datetime = timezone.make_aware(datetime.combine(to_date, datetime.max.time()))
            batch_ids_in_range = list(
                TotalStockModel.objects.filter(
                    created_at__range=(from_datetime, to_datetime)
                ).values_list('batch_id__batch_id', flat=True)
            )
            completed_batches = ModelMasterCreation.objects.filter(
                total_batch_quantity__gt=0,
                Moved_to_D_Picker=True,
                batch_id__in=batch_ids_in_range
            ).select_related('location', 'version').annotate(
                next_process_module=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('next_process_module')[:1])
            )
            report_data_completed = []
            for idx, batch in enumerate(completed_batches, start=1):
                lot_status = "Released"
                tray_cate_capacity = f"{batch.tray_type} - {batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ""
                remarks_holding = f"Holding Reason: {batch.holding_reason}" if batch.holding_reason else ""
                remarks_releasing = f"Releasing Reason: {batch.release_reason}" if batch.release_reason else ""
                combined_remarks = "\n".join(filter(None, [remarks_holding, remarks_releasing])) or ""
                row = {
                    'S.No': idx,
                    'Last Updated': batch.date_time.replace(tzinfo=None) if batch.date_time else None,
                    'Plating Stk No': batch.plating_stk_no or '',
                    'Polishing Stk No': batch.polishing_stk_no or '',
                    'Plating Color': batch.plating_color or '',
                    'Category': batch.category or '',
                    'Polish Finish': batch.polish_finish or '',
                    'Version': batch.version.version_name if batch.version else '',
                    'Tray Cate-Capacity': tray_cate_capacity,
                    'Source': f"{batch.vendor_internal}_{batch.location.location_name if batch.location else ''}",
                    'No of Tray': batch.no_of_trays or 0,
                    'Input Qty': batch.total_batch_quantity or 0,
                    'Process Status': 'T',
                    'Lot Status': lot_status,
                    'Current Stage': batch.next_process_module or '',
                    'Remarks': combined_remarks,
                }
                report_data_completed.append(row)
            df_completed = pd.DataFrame(report_data_completed)
            df_completed.to_excel(writer, sheet_name='Completed Table', index=False)

        elif module == 'input-screening':
            # Import required models and functions
            from modelmasterapp.models import ModelMasterCreation
            from InputScreening.models import IP_Rejection_ReasonStore
            from django.db.models import Q, F, Exists
            from django.utils import timezone
            import math
            from django.templatetags.static import static
            
            # Pick Table - Use same logic as IS_PickTable
            accepted_Ip_stock_subquery = TotalStockModel.objects.filter(
                batch_id=OuterRef('pk')
            ).values('accepted_Ip_stock')[:1]
            
            accepted_tray_scan_status_subquery = TotalStockModel.objects.filter(
                batch_id=OuterRef('pk')
            ).values('accepted_tray_scan_status')[:1]
            
            rejected_ip_stock_subquery = TotalStockModel.objects.filter(
                batch_id=OuterRef('pk')
            ).values('rejected_ip_stock')[:1]
            
            few_cases_accepted_Ip_stock_subquery = TotalStockModel.objects.filter(
                batch_id=OuterRef('pk')
            ).values('few_cases_accepted_Ip_stock')[:1]
            
            ip_onhold_picking_subquery = TotalStockModel.objects.filter(
                batch_id=OuterRef('pk')
            ).values('ip_onhold_picking')[:1]
            
            tray_verify_subquery = TotalStockModel.objects.filter(
                batch_id=OuterRef('pk')
            ).values('tray_verify')[:1]
            
            draft_tray_verify_subquery = TotalStockModel.objects.filter(
                batch_id=OuterRef('pk')
            ).values('draft_tray_verify')[:1]

            tray_scan_exists = Exists(
                TotalStockModel.objects.filter(
                    batch_id=OuterRef('pk'),
                    tray_scan_status=True
                )
            )

            pick_queryset = ModelMasterCreation.objects.filter(
                total_batch_quantity__gt=0,
            ).annotate(
                last_process_module=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('last_process_module')[:1]
                ),
                next_process_module=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('next_process_module')[:1]
                ),
                wiping_required=F('model_stock_no__wiping_required'),
                stock_lot_id=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('lot_id')[:1]
                ),
                ip_person_qty_verified=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('ip_person_qty_verified')[:1]
                ),
                lot_rejected_comment=Subquery(
                    IP_Rejection_ReasonStore.objects.filter(lot_id=OuterRef('stock_lot_id')).values('lot_rejected_comment')[:1]
                ),
                accepted_Ip_stock=accepted_Ip_stock_subquery,
                accepted_tray_scan_status=accepted_tray_scan_status_subquery,
                rejected_ip_stock=rejected_ip_stock_subquery,
                few_cases_accepted_Ip_stock=few_cases_accepted_Ip_stock_subquery,
                ip_onhold_picking=ip_onhold_picking_subquery,
                tray_verify=tray_verify_subquery,
                draft_tray_verify=draft_tray_verify_subquery,
                tray_scan_exists=tray_scan_exists,
                IP_pick_remarks=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('IP_pick_remarks')[:1]
                ),
                created_at=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('created_at')[:1]
                ),
                total_ip_accepted_quantity=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('total_IP_accpeted_quantity')[:1]
                ),
                ip_hold_lot=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('ip_hold_lot')[:1]
                ),
                ip_holding_reason=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('ip_holding_reason')[:1]
                ),
                ip_release_lot=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('ip_release_lot')[:1]
                ),
                ip_release_reason=Subquery(
                    TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('ip_release_reason')[:1]
                ),
                total_rejection_quantity=Subquery(
                    IP_Rejection_ReasonStore.objects.filter(lot_id=OuterRef('stock_lot_id')).values('total_rejection_quantity')[:1]
                ),
            ).filter(
                (Q(accepted_Ip_stock=False) | Q(accepted_Ip_stock__isnull=True)) &
                (Q(rejected_ip_stock=False) | Q(rejected_ip_stock__isnull=True)) &
                (Q(accepted_tray_scan_status=False) | Q(accepted_tray_scan_status__isnull=True)),
                tray_scan_exists=True,
            ).exclude(
                totalstockmodel__remove_lot=True    
            ).order_by('-date_time')

            pick_report_data = []
            for idx, batch in enumerate(pick_queryset, start=1):
                # Determine status based on the same logic as HTML table
                if batch.ip_hold_lot:
                    lot_status = "On Hold"
                elif batch.ip_release_lot:
                    lot_status = "Released"
                else:
                    lot_status = "In Process"
                
                # Calculate derived values
                tray_cate_capacity = f"{batch.tray_type} - {batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ""
                input_source = f"{batch.vendor_internal}_{batch.location.location_name if batch.location else ''}"
                ipa_wiping = "Yes" if batch.wiping_required else "No"
                no_of_trays = math.ceil(batch.total_batch_quantity / batch.tray_capacity) if batch.tray_capacity and batch.tray_capacity > 0 else 0
                lot_qty = batch.total_batch_quantity or 0
                
                # Calculate Accept Qty and Reject Qty
                accept_qty = batch.total_ip_accepted_quantity or 0
                reject_qty = batch.total_rejection_quantity or 0
                
                # Process Status and Current Stage
                process_status = "Draft"  # Default for pick table
                current_stage = batch.next_process_module or 'Input Screening'
                
                # Format Last Updated
                last_updated = ""
                if batch.date_time:
                    dt = batch.date_time.replace(tzinfo=None)
                    last_updated = dt.strftime('%d-%b-%y %I:%M %p')
                
                row = {
                    'S.No': idx,
                    'Last Updated': last_updated,
                    'Plating Stk No': batch.plating_stk_no or '',
                    'Polishing Stk No': batch.polishing_stk_no or '',
                    'Plating Color': batch.plating_color or '',
                    'Category': batch.category or '',
                    'Polish Finish': batch.polish_finish or '',
                    'Tray Cate-Capacity': tray_cate_capacity,
                    'Input Source': input_source,
                    'IPA Wiping': ipa_wiping,
                    'No of Trays': no_of_trays,
                    'LOT Qty': lot_qty,
                    'Accept Qty': accept_qty,
                    'Reject Qty': reject_qty,
                    'Process Status': process_status,
                    'Lot Status': lot_status,
                    'Current Stage': current_stage,
                    'Remarks': batch.IP_pick_remarks or '',
                }
                pick_report_data.append(row)
            
            df_pick = pd.DataFrame(pick_report_data)
            df_pick.to_excel(writer, sheet_name='Pick Table', index=False)
            
            # Completed Table - Use same logic as IS_Completed_Table
            tz = pytz.timezone("Asia/Kolkata")
            now_local = timezone.now().astimezone(tz)
            today = now_local.date()
            yesterday = today - timedelta(days=1)
            from_date = yesterday
            to_date = today
            from_datetime = timezone.make_aware(datetime.combine(from_date, datetime.min.time()))
            to_datetime = timezone.make_aware(datetime.combine(to_date, datetime.max.time()))

            completed_queryset = TotalStockModel.objects.filter(
                Q(accepted_Ip_stock=True) |
                Q(rejected_ip_stock=True) |
                (Q(few_cases_accepted_Ip_stock=True) & Q(ip_onhold_picking=False)),
                batch_id__total_batch_quantity__gt=0,
                remove_lot=False
            ).select_related(
                'batch_id',
                'batch_id__model_stock_no',
                'batch_id__version',
                'batch_id__location'
            ).annotate(
                batch_model_no=F('batch_id__model_stock_no__model_no'),
                batch_plating_color=F('batch_id__plating_color'),
                batch_polish_finish=F('batch_id__polish_finish'),
                batch_version_name=F('batch_id__version__version_name'),
                batch_version_internal=F('batch_id__version__version_internal'),
                batch_vendor_internal=F('batch_id__vendor_internal'),
                batch_location_name=F('batch_id__location__location_name'),
                batch_tray_type=F('batch_id__tray_type'),
                batch_tray_capacity=F('batch_id__tray_capacity'),
                batch_moved_to_d_picker=F('batch_id__Moved_to_D_Picker'),
                batch_draft_saved=F('batch_id__Draft_Saved'),
                batch_total_batch_quantity=F('batch_id__total_batch_quantity'),
                batch_date_time=F('batch_id__date_time'),
                batch_plating_stk_no=F('batch_id__plating_stk_no'),
                batch_polishing_stk_no=F('batch_id__polishing_stk_no'),
                batch_category=F('batch_id__category')
            ).order_by('-last_process_date_time')

            completed_report_data = []
            for idx, total_stock_obj in enumerate(completed_queryset, start=1):
                # Calculate no_of_trays
                no_of_trays = math.ceil(total_stock_obj.total_stock / total_stock_obj.batch_tray_capacity) if total_stock_obj.batch_tray_capacity and total_stock_obj.batch_tray_capacity > 0 else 0
                
                # Determine status
                if total_stock_obj.accepted_Ip_stock:
                    status = "Accepted"
                elif total_stock_obj.rejected_ip_stock:
                    status = "Rejected"
                elif total_stock_obj.few_cases_accepted_Ip_stock and not total_stock_obj.ip_onhold_picking:
                    status = "Partially Accepted"
                else:
                    status = "Completed"
                
                row = {
                    'S.No': idx,
                    'Date & Time': total_stock_obj.batch_date_time.replace(tzinfo=None) if total_stock_obj.batch_date_time else None,
                    'Plating Stock No': total_stock_obj.batch_plating_stk_no or '',
                    'Polishing Stock No': total_stock_obj.batch_polishing_stk_no or '',
                    'Model No': total_stock_obj.batch_model_no or '',
                    'Plating Color': total_stock_obj.batch_plating_color or '',
                    'Polish Finish': total_stock_obj.batch_polish_finish or '',
                    'Version': total_stock_obj.batch_version_name or '',
                    'Vendor': total_stock_obj.batch_vendor_internal or '',
                    'Location': total_stock_obj.batch_location_name or '',
                    'Tray Type': total_stock_obj.batch_tray_type or '',
                    'Tray Capacity': total_stock_obj.batch_tray_capacity or 0,
                    'No of Trays': no_of_trays,
                    'Total Stock': total_stock_obj.total_stock or 0,
                    'Lot ID': total_stock_obj.lot_id or '',
                    'Status': status,
                    'IP Pick Remarks': total_stock_obj.IP_pick_remarks or '',
                    'Last Process Date Time': total_stock_obj.last_process_date_time.replace(tzinfo=None) if total_stock_obj.last_process_date_time else None,
                }
                completed_report_data.append(row)
            
            df_completed = pd.DataFrame(completed_report_data)
            df_completed.to_excel(writer, sheet_name='Completed Table', index=False)
            
            # Accept Table - Use same logic as IS_AcceptTable
            accept_queryset = TotalStockModel.objects.filter(
                Q(accepted_Ip_stock=True) |  
                (Q(few_cases_accepted_Ip_stock=True) & Q(ip_onhold_picking=False)) |  
                (Q(few_cases_accepted_Ip_stock=True) & Q(ip_onhold_picking=True)),  
                batch_id__total_batch_quantity__gt=0,
                remove_lot=False,  
            ).select_related(
                'batch_id',
                'batch_id__model_stock_no',
                'batch_id__version',
                'batch_id__location'
            ).annotate(
                batch_model_no=F('batch_id__model_stock_no__model_no'),
                batch_plating_color=F('batch_id__plating_color'),
                batch_polish_finish=F('batch_id__polish_finish'),
                batch_version_name=F('batch_id__version__version_name'),
                batch_version_internal=F('batch_id__version__version_internal'),
                batch_vendor_internal=F('batch_id__vendor_internal'),
                batch_location_name=F('batch_id__location__location_name'),
                batch_tray_type=F('batch_id__tray_type'),
                batch_tray_capacity=F('batch_id__tray_capacity'),
                batch_moved_to_d_picker=F('batch_id__Moved_to_D_Picker'),
                batch_draft_saved=F('batch_id__Draft_Saved'),
                batch_total_batch_quantity=F('batch_id__total_batch_quantity'),
                batch_date_time=F('batch_id__date_time'),
                batch_plating_stk_no=F('batch_id__plating_stk_no'),
                batch_polishing_stk_no=F('batch_id__polishing_stk_no'),
                batch_category=F('batch_id__category')
            ).order_by('-last_process_date_time')

            accept_report_data = []
            for idx, total_stock_obj in enumerate(accept_queryset, start=1):
                # Calculate accepted quantity (same logic as IS_AcceptTable)
                total_IP_accpeted_quantity = total_stock_obj.total_IP_accpeted_quantity
                lot_id = total_stock_obj.lot_id
                total_stock = total_stock_obj.total_stock
                
                if total_IP_accpeted_quantity and total_IP_accpeted_quantity > 0:
                    display_accepted_qty = total_IP_accpeted_quantity
                else:
                    total_rejection_qty = 0
                    rejection_store = IP_Rejection_ReasonStore.objects.filter(lot_id=lot_id).first()
                    if rejection_store and rejection_store.total_rejection_quantity:
                        total_rejection_qty = rejection_store.total_rejection_quantity
                    
                    if total_stock > 0 and total_rejection_qty > 0:
                        display_accepted_qty = max(total_stock - total_rejection_qty, 0)
                    else:
                        display_accepted_qty = 0
                
                # Calculate no_of_trays
                no_of_trays = math.ceil(display_accepted_qty / total_stock_obj.batch_tray_capacity) if total_stock_obj.batch_tray_capacity and total_stock_obj.batch_tray_capacity > 0 else 0
                
                row = {
                    'S.No': idx,
                    'Date & Time': total_stock_obj.batch_date_time.replace(tzinfo=None) if total_stock_obj.batch_date_time else None,
                    'Plating Stock No': total_stock_obj.batch_plating_stk_no or '',
                    'Polishing Stock No': total_stock_obj.batch_polishing_stk_no or '',
                    'Model No': total_stock_obj.batch_model_no or '',
                    'Plating Color': total_stock_obj.batch_plating_color or '',
                    'Polish Finish': total_stock_obj.batch_polish_finish or '',
                    'Version': total_stock_obj.batch_version_name or '',
                    'Vendor': total_stock_obj.batch_vendor_internal or '',
                    'Location': total_stock_obj.batch_location_name or '',
                    'Tray Type': total_stock_obj.batch_tray_type or '',
                    'Tray Capacity': total_stock_obj.batch_tray_capacity or 0,
                    'No of Trays': no_of_trays,
                    'Accept Qty': display_accepted_qty,
                    'Lot ID': total_stock_obj.lot_id or '',
                    'IP Pick Remarks': total_stock_obj.IP_pick_remarks or '',
                    'Last Process Date Time': total_stock_obj.last_process_date_time.replace(tzinfo=None) if total_stock_obj.last_process_date_time else None,
                }
                accept_report_data.append(row)
            
            df_accept = pd.DataFrame(accept_report_data)
            df_accept.to_excel(writer, sheet_name='Accept Table', index=False)
            
            # Reject Table - Use same logic as IS_RejectTable
            reject_queryset = TotalStockModel.objects.filter(
                Q(rejected_ip_stock=True) |  
                (Q(few_cases_accepted_Ip_stock=True) & Q(ip_onhold_picking=False)) |  
                (Q(few_cases_accepted_Ip_stock=True) & Q(ip_onhold_picking=True)),  
                batch_id__total_batch_quantity__gt=0,
                remove_lot=False,  
            ).select_related(
                'batch_id',
                'batch_id__model_stock_no',
                'batch_id__version',
                'batch_id__location'
            ).annotate(
                batch_model_no=F('batch_id__model_stock_no__model_no'),
                batch_plating_color=F('batch_id__plating_color'),
                batch_polish_finish=F('batch_id__polish_finish'),
                batch_version_name=F('batch_id__version__version_name'),
                batch_version_internal=F('batch_id__version__version_internal'),
                batch_vendor_internal=F('batch_id__vendor_internal'),
                batch_location_name=F('batch_id__location__location_name'),
                batch_tray_type=F('batch_id__tray_type'),
                batch_tray_capacity=F('batch_id__tray_capacity'),
                batch_moved_to_d_picker=F('batch_id__Moved_to_D_Picker'),
                batch_draft_saved=F('batch_id__Draft_Saved'),
                batch_total_batch_quantity=F('batch_id__total_batch_quantity'),
                batch_date_time=F('batch_id__date_time'),
                batch_plating_stk_no=F('batch_id__plating_stk_no'),
                batch_polishing_stk_no=F('batch_id__polishing_stk_no'),
                batch_category=F('batch_id__category')
            ).order_by('-last_process_date_time')

            reject_report_data = []
            for idx, total_stock_obj in enumerate(reject_queryset, start=1):
                stock_lot_id = total_stock_obj.lot_id
                
                # Get rejection quantity (same logic as IS_RejectTable)
                rejection_qty = 0
                rejection_record = IP_Rejection_ReasonStore.objects.filter(lot_id=stock_lot_id).first()
                if rejection_record and rejection_record.total_rejection_quantity:
                    rejection_qty = rejection_record.total_rejection_quantity
                
                # Get rejection reasons
                rejection_letters = []
                batch_rejection = False
                lot_rejected_comment = None
                
                if rejection_record:
                    batch_rejection = rejection_record.batch_rejection
                    lot_rejected_comment = rejection_record.lot_rejected_comment
                    reasons = rejection_record.rejection_reason.all()
                    for r in reasons:
                        if r.rejection_reason.upper() != 'SHORTAGE':
                            rejection_letters.append(r.rejection_reason[0].upper())
                
                # Check for SHORTAGE
                shortage_exists = IP_Rejected_TrayScan.objects.filter(
                    lot_id=stock_lot_id,
                    rejection_reason__rejection_reason__iexact='SHORTAGE'
                ).exists()
                if shortage_exists:
                    rejection_letters.append('S')
                
                # Get shortage quantity
                shortage_qty = sum(
                    int(obj.rejected_tray_quantity or 0)
                    for obj in IP_Rejected_TrayScan.objects.filter(
                        lot_id=stock_lot_id,
                        rejection_reason__rejection_reason__iexact='SHORTAGE'
                    )
                )
                
                # Calculate effective stock and trays
                effective_stock = max(rejection_qty - shortage_qty, 0)
                no_of_trays = math.ceil(effective_stock / total_stock_obj.batch_tray_capacity) if total_stock_obj.batch_tray_capacity and total_stock_obj.batch_tray_capacity > 0 else 0
                
                row = {
                    'S.No': idx,
                    'Date & Time': total_stock_obj.batch_date_time.replace(tzinfo=None) if total_stock_obj.batch_date_time else None,
                    'Plating Stock No': total_stock_obj.batch_plating_stk_no or '',
                    'Polishing Stock No': total_stock_obj.batch_polishing_stk_no or '',
                    'Model No': total_stock_obj.batch_model_no or '',
                    'Plating Color': total_stock_obj.batch_plating_color or '',
                    'Polish Finish': total_stock_obj.batch_polish_finish or '',
                    'Version': total_stock_obj.batch_version_name or '',
                    'Vendor': total_stock_obj.batch_vendor_internal or '',
                    'Location': total_stock_obj.batch_location_name or '',
                    'Tray Type': total_stock_obj.batch_tray_type or '',
                    'Tray Capacity': total_stock_obj.batch_tray_capacity or 0,
                    'No of Trays': no_of_trays,
                    'Reject Qty': rejection_qty,
                    'Shortage Qty': shortage_qty,
                    'Rejection Reasons': ','.join(rejection_letters),
                    'Batch Rejection': 'Yes' if batch_rejection else 'No',
                    'Lot Rejected Comment': lot_rejected_comment or '',
                    'Lot ID': total_stock_obj.lot_id or '',
                    'IP Pick Remarks': total_stock_obj.IP_pick_remarks or '',
                    'Last Process Date Time': total_stock_obj.last_process_date_time.replace(tzinfo=None) if total_stock_obj.last_process_date_time else None,
                }
                reject_report_data.append(row)
            
            df_reject = pd.DataFrame(reject_report_data)
            df_reject.to_excel(writer, sheet_name='Reject Table', index=False)

        elif module == 'brass-qc':
            # Import required models and functions
            from modelmasterapp.models import ModelMasterCreation
            from InputScreening.models import IP_Rejection_ReasonStore
            from django.db.models import Q, F, Exists
            from django.utils import timezone
            import math
            from django.templatetags.static import static
            
            # Pick Table - Use same logic as BrassPickTableView
            brass_rejection_qty_subquery = Brass_QC_Rejection_ReasonStore.objects.filter(
                lot_id=OuterRef('lot_id')
            ).values('total_rejection_quantity')[:1]

            has_draft_subquery = Exists(
                Brass_QC_Draft_Store.objects.filter(
                    lot_id=OuterRef('lot_id')
                )
            )
            
            draft_type_subquery = Brass_QC_Draft_Store.objects.filter(
                lot_id=OuterRef('lot_id')
            ).values('draft_type')[:1]

            pick_queryset = TotalStockModel.objects.select_related(
                'batch_id',
                'batch_id__model_stock_no',
                'batch_id__version',
                'batch_id__location'
            ).filter(
                batch_id__total_batch_quantity__gt=0
            ).annotate(
                wiping_required=F('batch_id__model_stock_no__wiping_required'),
                has_draft=has_draft_subquery,
                draft_type=draft_type_subquery,
                brass_rejection_total_qty=brass_rejection_qty_subquery,
            ).filter(
                (
                    (
                        Q(brass_qc_accptance__isnull=True) | Q(brass_qc_accptance=False)
                    ) &
                    (
                        Q(brass_qc_rejection__isnull=True) | Q(brass_qc_rejection=False)
                    ) &
                    ~Q(brass_qc_few_cases_accptance=True, brass_onhold_picking=False)
                    &
                    (
                        Q(accepted_Ip_stock=True) | 
                        Q(few_cases_accepted_Ip_stock=True, ip_onhold_picking=False)
                    )
                )
                |
                Q(send_brass_qc=True)
                |
                Q(brass_qc_rejection=True, brass_onhold_picking=True)
                |
                Q(send_brass_audit_to_qc=True)
            ).exclude(
                Q(brass_audit_rejection=True)
            ).order_by('-last_process_date_time', '-lot_id')

            pick_report_data = []
            for idx, stock_obj in enumerate(pick_queryset, start=1):
                batch = stock_obj.batch_id
                
                # Calculate derived values
                tray_cate_capacity = f"{batch.tray_type} - {batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ""
                input_source = f"{batch.vendor_internal}_{batch.location.location_name if batch.location else ''}"
                
                # Lot Qty
                lot_qty = batch.total_batch_quantity or 0
                
                # Physical Qty
                physical_qty = stock_obj.brass_physical_qty or 0
                
                # Accept Qty calculation
                total_IP_accpeted_quantity = stock_obj.total_IP_accpeted_quantity or 0
                if total_IP_accpeted_quantity > 0:
                    accept_qty = total_IP_accpeted_quantity
                else:
                    total_rejection_qty = 0
                    rejection_store = IP_Rejection_ReasonStore.objects.filter(lot_id=stock_obj.lot_id).first()
                    if rejection_store and rejection_store.total_rejection_quantity:
                        total_rejection_qty = rejection_store.total_rejection_quantity
                    
                    if stock_obj.total_stock > 0 and total_rejection_qty > 0:
                        accept_qty = max(stock_obj.total_stock - total_rejection_qty, 0)
                    else:
                        accept_qty = 0
                
                # Lot Qty (same as Accept Qty in pick table)
                lot_qty = accept_qty
                
                # Reject Qty
                reject_qty = stock_obj.brass_rejection_total_qty or 0
                
                # No of Trays
                no_of_trays = math.ceil(accept_qty / batch.tray_capacity) if batch.tray_capacity and batch.tray_capacity > 0 and accept_qty > 0 else 0
                
                # Process Status
                process_status = "Draft"  # Default for pick table
                
                # Lot Status
                if stock_obj.brass_hold_lot:
                    lot_status = "On Hold"
                elif stock_obj.brass_release_lot:
                    lot_status = "Released"
                else:
                    lot_status = "In Process"
                
                # Current Stage
                current_stage = stock_obj.next_process_module or 'Brass QC'
                
                # Last Updated
                last_updated = ""
                if stock_obj.last_process_date_time:
                    dt = stock_obj.last_process_date_time.replace(tzinfo=None)
                    last_updated = dt.strftime('%d-%b-%y %I:%M %p')
                
                row = {
                    'S.No': idx,
                    'Last Updated': last_updated,
                    'Plating Stk No': batch.plating_stk_no or '',
                    'Polishing Stk No': batch.polishing_stk_no or '',
                    'Plating Color': batch.plating_color or '',
                    'Category': batch.category or '',
                    'Polish Finish': batch.polish_finish or '',
                    'Tray Cate-Capacity': tray_cate_capacity,
                    'Input Source': input_source,
                    'No of Trays': no_of_trays,
                    'Lot Qty': lot_qty,
                    'Physical Qty': physical_qty,
                    'Accept Qty': accept_qty,
                    'Reject Qty': reject_qty,
                    'Process Status': process_status,
                    'Lot Status': lot_status,
                    'Current Stage': current_stage,
                    'Remarks': stock_obj.Bq_pick_remarks or '',
                }
                pick_report_data.append(row)
            
            df_pick = pd.DataFrame(pick_report_data)
            df_pick.to_excel(writer, sheet_name='Pick Table', index=False)
            
            # Completed Table - Use same logic as BrassCompletedView
            tz = pytz.timezone("Asia/Kolkata")
            now_local = timezone.now().astimezone(tz)
            today = now_local.date()
            yesterday = today - timedelta(days=1)
            from_date = yesterday
            to_date = today
            from_datetime = timezone.make_aware(datetime.combine(from_date, datetime.min.time()))
            to_datetime = timezone.make_aware(datetime.combine(to_date, datetime.max.time()))

            brass_rejection_qty_subquery = Brass_QC_Rejection_ReasonStore.objects.filter(
                lot_id=OuterRef('lot_id')
            ).values('total_rejection_quantity')[:1]

            completed_queryset = TotalStockModel.objects.select_related(
                'batch_id',
                'batch_id__model_stock_no',
                'batch_id__version',
                'batch_id__location'
            ).filter(
                batch_id__total_batch_quantity__gt=0,
                bq_last_process_date_time__range=(from_datetime, to_datetime)
            ).annotate(
                brass_rejection_qty=brass_rejection_qty_subquery,
            ).filter(
                Q(brass_qc_accptance=True) |
                Q(brass_qc_rejection=True) |
                Q(brass_qc_few_cases_accptance=True, brass_onhold_picking=False)
            ).order_by('-bq_last_process_date_time', '-lot_id')

            completed_report_data = []
            for idx, stock_obj in enumerate(completed_queryset, start=1):
                batch = stock_obj.batch_id
                
                # Calculate derived values
                tray_cate_capacity = f"{batch.tray_type} - {batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ""
                input_source = f"{batch.vendor_internal}_{batch.location.location_name if batch.location else ''}"
                
                # Physical Qty
                physical_qty = stock_obj.brass_physical_qty or 0
                
                # Accept Qty calculation
                total_IP_accpeted_quantity = stock_obj.total_IP_accpeted_quantity or 0
                if total_IP_accpeted_quantity > 0:
                    accept_qty = total_IP_accpeted_quantity
                else:
                    total_rejection_qty = 0
                    rejection_store = IP_Rejection_ReasonStore.objects.filter(lot_id=stock_obj.lot_id).first()
                    if rejection_store and rejection_store.total_rejection_quantity:
                        total_rejection_qty = rejection_store.total_rejection_quantity
                    
                    if stock_obj.total_stock > 0 and total_rejection_qty > 0:
                        accept_qty = max(stock_obj.total_stock - total_rejection_qty, 0)
                    else:
                        accept_qty = 0
                
                # Lot Qty
                lot_qty = accept_qty
                
                # Reject Qty
                reject_qty = stock_obj.brass_rejection_qty or 0
                
                # No of Trays
                no_of_trays = math.ceil(accept_qty / batch.tray_capacity) if batch.tray_capacity and batch.tray_capacity > 0 and accept_qty > 0 else 0
                
                # Process Status
                if stock_obj.brass_qc_accptance:
                    process_status = "Accepted"
                elif stock_obj.brass_qc_rejection:
                    process_status = "Rejected"
                elif stock_obj.brass_qc_few_cases_accptance and not stock_obj.brass_onhold_picking:
                    process_status = "Partially Accepted"
                else:
                    process_status = "Completed"
                
                # Lot Status
                if stock_obj.brass_hold_lot:
                    lot_status = "On Hold"
                elif stock_obj.brass_release_lot:
                    lot_status = "Released"
                else:
                    lot_status = "Completed"
                
                # Current Stage
                current_stage = stock_obj.next_process_module or ''
                
                # Last Updated
                last_updated = ""
                if stock_obj.bq_last_process_date_time:
                    dt = stock_obj.bq_last_process_date_time.replace(tzinfo=None)
                    last_updated = dt.strftime('%d-%b-%y %I:%M %p')
                
                row = {
                    'S.No': idx,
                    'Last Updated': last_updated,
                    'Plating Stk No': batch.plating_stk_no or '',
                    'Polishing Stk No': batch.polishing_stk_no or '',
                    'Plating Color': batch.plating_color or '',
                    'Category': batch.category or '',
                    'Polish Finish': batch.polish_finish or '',
                    'Tray Cate-Capacity': tray_cate_capacity,
                    'Input Source': input_source,
                    'No of Trays': no_of_trays,
                    'Lot Qty': lot_qty,
                    'Physical Qty': physical_qty,
                    'Accept Qty': accept_qty,
                    'Reject Qty': reject_qty,
                    'Process Status': process_status,
                    'Lot Status': lot_status,
                    'Current Stage': current_stage,
                    'Remarks': stock_obj.Bq_pick_remarks or '',
                }
                completed_report_data.append(row)
            
            df_completed = pd.DataFrame(completed_report_data)
            df_completed.to_excel(writer, sheet_name='Completed Table', index=False)

        elif module == 'iqf':
            from modelmasterapp.models import ModelMasterCreation
            # Pick Table
            pick_batches = ModelMasterCreation.objects.select_related('version', 'location').annotate(
                send_brass_audit_to_iqf=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('send_brass_audit_to_iqf')[:1]),
                iqf_hold_lot=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_hold_lot')[:1]),
                iqf_few_cases_acceptance=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_few_cases_acceptance')[:1])
            ).filter(send_brass_audit_to_iqf=True).exclude(lot_id__in=IQF_Accepted_TrayID_Store.objects.values_list('lot_id', flat=True))
            report_data = []
            for idx, batch in enumerate(pick_batches, start=1):
                try:
                    lot_id = batch.lot_id
                    accepted_scans = IQF_Accepted_TrayScan.objects.filter(lot_id=lot_id).values_list('accepted_tray_quantity', flat=True)
                    accept_qty = sum(int(qty) for qty in accepted_scans if qty and qty.isdigit()) or 0
                    rejected_scans = IQF_Rejected_TrayScan.objects.filter(lot_id=lot_id).values_list('rejected_tray_quantity', flat=True)
                    reject_qty = sum(int(qty) for qty in rejected_scans if qty and qty.isdigit()) or 0
                    last_updated = batch.date_time
                    lot_status = 'Hold' if batch.iqf_hold_lot else 'Active'
                    remarks = batch.holding_reason or batch.release_reason or ''
                    tray_cate_capacity = f"{batch.tray_type}-{batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ''
                    info = []
                    if batch.send_brass_audit_to_iqf: info.append("From Brass Audit")
                    if batch.iqf_hold_lot: info.append("Hold")
                    if batch.iqf_few_cases_acceptance: info.append("Few Cases Acceptance")
                    info_str = ', '.join(info)
                    no_of_trays = IQFTrayId.objects.filter(lot_id=lot_id).count()
                    physical_qty = IQFTrayId.objects.filter(lot_id=lot_id).aggregate(Sum('tray_quantity'))['tray_quantity__sum'] or 0
                    input_source = batch.location.location_name if batch.location else ''
                    row = {
                        'S.No': idx,
                        'Info': info_str,
                        'Last Updated': last_updated,
                        'Plating Stk No': batch.plating_stk_no or '',
                        'Polishing Stk No': batch.polishing_stk_no or '',
                        'Plating Color': batch.plating_color or '',
                        'Category': batch.category or '',
                        'Polish Finish': batch.polish_finish or '',
                        'Tray Cate-Capacity': tray_cate_capacity,
                        'Input Source': input_source,
                        'No of Trays': no_of_trays,
                        'RW Qty': batch.total_batch_quantity or 0,
                        'Physical Qty': physical_qty,
                        'Accept Qty': accept_qty,
                        'Reject Qty': reject_qty,
                        'Process Status': 'Pick',
                        'Action': '',
                    }
                    report_data.append(row)
                except Exception as e:
                    continue
            report_data = convert_datetimes(report_data)
            df = pd.DataFrame(report_data)
            if not df.empty:
                df.to_excel(writer, sheet_name='Pick Table', index=False)
            else:
                empty_df = pd.DataFrame(columns=['S.No', 'Info', 'Last Updated', 'Plating Stk No', 'Polishing Stk No', 'Plating Color', 'Category', 'Polish Finish', 'Tray Cate-Capacity', 'Input Source', 'No of Trays', 'RW Qty', 'Physical Qty', 'Accept Qty', 'Reject Qty', 'Process Status', 'Action'])
                empty_df.to_excel(writer, sheet_name='Pick Table', index=False)

            # Completed Table
            completed_lot_ids = IQF_Accepted_TrayID_Store.objects.values_list('lot_id', flat=True).distinct()
            completed_batches = ModelMasterCreation.objects.filter(lot_id__in=completed_lot_ids).select_related('version', 'location').annotate(
                iqf_hold_lot=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_hold_lot')[:1]),
                iqf_few_cases_acceptance=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_few_cases_acceptance')[:1])
            )
            report_data = []
            for idx, batch in enumerate(completed_batches, start=1):
                try:
                    lot_id = batch.lot_id
                    accepted_scans = IQF_Accepted_TrayScan.objects.filter(lot_id=lot_id).values_list('accepted_tray_quantity', flat=True)
                    accept_qty = sum(int(qty) for qty in accepted_scans if qty and qty.isdigit()) or 0
                    rejected_scans = IQF_Rejected_TrayScan.objects.filter(lot_id=lot_id).values_list('rejected_tray_quantity', flat=True)
                    reject_qty = sum(int(qty) for qty in rejected_scans if qty and qty.isdigit()) or 0
                    last_updated = batch.date_time
                    lot_status = 'Hold' if batch.iqf_hold_lot else 'Active'
                    remarks = batch.holding_reason or batch.release_reason or ''
                    tray_cate_capacity = f"{batch.tray_type}-{batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ''
                    info = []
                    if batch.iqf_hold_lot: info.append("Hold")
                    if batch.iqf_few_cases_acceptance: info.append("Few Cases Acceptance")
                    info_str = ', '.join(info)
                    no_of_trays = IQFTrayId.objects.filter(lot_id=lot_id).count()
                    physical_qty = IQFTrayId.objects.filter(lot_id=lot_id).aggregate(Sum('tray_quantity'))['tray_quantity__sum'] or 0
                    input_source = batch.location.location_name if batch.location else ''
                    row = {
                        'S.No': idx,
                        'Info': info_str,
                        'Last Updated': last_updated,
                        'Plating Stk No': batch.plating_stk_no or '',
                        'Polishing Stk No': batch.polishing_stk_no or '',
                        'Plating Color': batch.plating_color or '',
                        'Category': batch.category or '',
                        'Polish Finish': batch.polish_finish or '',
                        'Tray Cate-Capacity': tray_cate_capacity,
                        'Input Source': input_source,
                        'No of Trays': no_of_trays,
                        'RW Qty': batch.total_batch_quantity or 0,
                        'Physical Qty': physical_qty,
                        'Accept Qty': accept_qty,
                        'Reject Qty': reject_qty,
                        'Process Status': 'Completed',
                        'Action': '',
                        'Lot Status': lot_status,
                        'Current Stage': 'IQF Completed',
                        'Remarks': remarks,
                    }
                    report_data.append(row)
                except Exception as e:
                    continue
            report_data = convert_datetimes(report_data)
            df = pd.DataFrame(report_data)
            if not df.empty:
                df.to_excel(writer, sheet_name='Completed Table', index=False)
            else:
                empty_df = pd.DataFrame(columns=['S.No', 'Info', 'Last Updated', 'Plating Stk No', 'Polishing Stk No', 'Plating Color', 'Category', 'Polish Finish', 'Tray Cate-Capacity', 'Input Source', 'No of Trays', 'RW Qty', 'Physical Qty', 'Accept Qty', 'Reject Qty', 'Process Status', 'Action', 'Lot Status', 'Current Stage', 'Remarks'])
                empty_df.to_excel(writer, sheet_name='Completed Table', index=False)

            # Accept Table
            accept_lot_ids = IQF_Accepted_TrayScan.objects.values_list('lot_id', flat=True).distinct()
            accept_batches = ModelMasterCreation.objects.filter(lot_id__in=accept_lot_ids).select_related('version', 'location').annotate(
                iqf_hold_lot=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_hold_lot')[:1]),
                iqf_few_cases_acceptance=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_few_cases_acceptance')[:1])
            )
            report_data = []
            for idx, batch in enumerate(accept_batches, start=1):
                try:
                    lot_id = batch.lot_id
                    accepted_scans = IQF_Accepted_TrayScan.objects.filter(lot_id=lot_id).values_list('accepted_tray_quantity', flat=True)
                    accept_qty = sum(int(qty) for qty in accepted_scans if qty and qty.isdigit()) or 0
                    rejected_scans = IQF_Rejected_TrayScan.objects.filter(lot_id=lot_id).values_list('rejected_tray_quantity', flat=True)
                    reject_qty = sum(int(qty) for qty in rejected_scans if qty and qty.isdigit()) or 0
                    last_updated = batch.date_time
                    tray_cate_capacity = f"{batch.tray_type}-{batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ''
                    info = []
                    if batch.iqf_hold_lot: info.append("Hold")
                    if batch.iqf_few_cases_acceptance: info.append("Few Cases Acceptance")
                    info_str = ', '.join(info)
                    no_of_trays = IQFTrayId.objects.filter(lot_id=lot_id).count()
                    physical_qty = IQFTrayId.objects.filter(lot_id=lot_id).aggregate(Sum('tray_quantity'))['tray_quantity__sum'] or 0
                    input_source = batch.location.location_name if batch.location else ''
                    row = {
                        'S.No': idx,
                        'Info': info_str,
                        'Last Updated': last_updated,
                        'Plating Stk No': batch.plating_stk_no or '',
                        'Polishing Stk No': batch.polishing_stk_no or '',
                        'Plating Color': batch.plating_color or '',
                        'Category': batch.category or '',
                        'Polish Finish': batch.polish_finish or '',
                        'Tray Cate-Capacity': tray_cate_capacity,
                        'Input Source': input_source,
                        'No of Trays': no_of_trays,
                        'RW Qty': batch.total_batch_quantity or 0,
                        'Physical Qty': physical_qty,
                        'Accept Qty': accept_qty,
                        'Reject Qty': reject_qty,
                        'Process Status': 'Accept',
                        'Action': '',
                    }
                    report_data.append(row)
                except Exception as e:
                    continue
            report_data = convert_datetimes(report_data)
            df = pd.DataFrame(report_data)
            if not df.empty:
                df.to_excel(writer, sheet_name='Accept Table', index=False)
            else:
                empty_df = pd.DataFrame(columns=['S.No', 'Info', 'Last Updated', 'Plating Stk No', 'Polishing Stk No', 'Plating Color', 'Category', 'Polish Finish', 'Tray Cate-Capacity', 'Input Source', 'No of Trays', 'RW Qty', 'Physical Qty', 'Accept Qty', 'Reject Qty', 'Process Status', 'Action'])
                empty_df.to_excel(writer, sheet_name='Accept Table', index=False)

            # Reject Table
            reject_lot_ids = IQF_Rejected_TrayScan.objects.values_list('lot_id', flat=True).distinct()
            reject_batches = ModelMasterCreation.objects.filter(lot_id__in=reject_lot_ids).select_related('version', 'location').annotate(
                iqf_hold_lot=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_hold_lot')[:1]),
                iqf_few_cases_acceptance=Subquery(TotalStockModel.objects.filter(batch_id=OuterRef('pk')).values('iqf_few_cases_acceptance')[:1])
            )
            report_data = []
            for idx, batch in enumerate(reject_batches, start=1):
                try:
                    lot_id = batch.lot_id
                    accepted_scans = IQF_Accepted_TrayScan.objects.filter(lot_id=lot_id).values_list('accepted_tray_quantity', flat=True)
                    accept_qty = sum(int(qty) for qty in accepted_scans if qty and qty.isdigit()) or 0
                    rejected_scans = IQF_Rejected_TrayScan.objects.filter(lot_id=lot_id).values_list('rejected_tray_quantity', flat=True)
                    reject_qty = sum(int(qty) for qty in rejected_scans if qty and qty.isdigit()) or 0
                    last_updated = batch.date_time
                    tray_cate_capacity = f"{batch.tray_type}-{batch.tray_capacity}" if batch.tray_type and batch.tray_capacity else ''
                    info = []
                    if batch.iqf_hold_lot: info.append("Hold")
                    if batch.iqf_few_cases_acceptance: info.append("Few Cases Acceptance")
                    info_str = ', '.join(info)
                    no_of_trays = IQFTrayId.objects.filter(lot_id=lot_id).count()
                    physical_qty = IQFTrayId.objects.filter(lot_id=lot_id).aggregate(Sum('tray_quantity'))['tray_quantity__sum'] or 0
                    reject_reasons = IQF_Rejection_ReasonStore.objects.filter(lot_id=lot_id).values_list('rejection_reason__rejection_reason', flat=True)
                    reject_reason_str = ', '.join(reject_reasons) if reject_reasons else ''
                    lot_remark = ''
                    input_source = batch.location.location_name if batch.location else ''
                    row = {
                        'S.No': idx,
                        'Info': info_str,
                        'Last Updated': last_updated,
                        'Plating Stk No': batch.plating_stk_no or '',
                        'Polishing Stk No': batch.polishing_stk_no or '',
                        'Plating Color': batch.plating_color or '',
                        'Polish Finish': batch.polish_finish or '',
                        'Source - Location': input_source,
                        'Tray Type Capacity': tray_cate_capacity,
                        'Tray Cate-Capacity': tray_cate_capacity,
                        'Reject Qty': reject_qty,
                        'Reject Reason': reject_reason_str,
                        'Lot Remark': lot_remark,
                    }
                    report_data.append(row)
                except Exception as e:
                    continue
            report_data = convert_datetimes(report_data)
            df = pd.DataFrame(report_data)
            if not df.empty:
                df.to_excel(writer, sheet_name='Reject Table', index=False)
            else:
                empty_df = pd.DataFrame(columns=['S.No', 'Info', 'Last Updated', 'Plating Stk No', 'Polishing Stk No', 'Plating Color', 'Polish Finish', 'Source - Location', 'Tray Type Capacity', 'Tray Cate-Capacity', 'Reject Qty', 'Reject Reason', 'Lot Remark'])
                empty_df.to_excel(writer, sheet_name='Reject Table', index=False)

        elif module == 'brass-audit':
            from BrassAudit.models import Brass_Audit_Rejection_Table, Brass_Audit_Rejection_ReasonStore, Brass_Audit_Draft_Store
            from django.db.models import Exists, F, Q
            from django.utils import timezone

            # Pick Table - Similar to BrassAuditPickTableView
            brass_rejection_qty_subquery = Brass_Audit_Rejection_ReasonStore.objects.filter(
                lot_id=OuterRef('lot_id')
            ).values('total_rejection_quantity')[:1]

            has_draft_subquery = Exists(
                Brass_Audit_Draft_Store.objects.filter(
                    lot_id=OuterRef('lot_id')
                )
            )
            
            draft_type_subquery = Brass_Audit_Draft_Store.objects.filter(
                lot_id=OuterRef('lot_id')
            ).values('draft_type')[:1]

            pick_queryset = TotalStockModel.objects.select_related(
                'batch_id',
                'batch_id__model_stock_no',
                'batch_id__version',
                'batch_id__location'
            ).filter(
                batch_id__total_batch_quantity__gt=0
            ).annotate(
                wiping_required=F('batch_id__model_stock_no__wiping_required'),
                has_draft=has_draft_subquery,
                draft_type=draft_type_subquery,
                brass_rejection_total_qty=brass_rejection_qty_subquery,
            ).filter(
                Q(brass_qc_accptance=True, brass_audit_accptance__isnull=True) |
                Q(brass_qc_accptance=True, brass_audit_accptance=False) |
                Q(brass_qc_few_cases_accptance=True, brass_onhold_picking=False)|
                Q(brass_audit_few_cases_accptance=True, brass_audit_onhold_picking=True)
            ).exclude(
                brass_audit_accptance=True
            ).exclude(
                brass_audit_rejection=True
            ).exclude(
                Q(brass_audit_few_cases_accptance=True, brass_audit_onhold_picking=False)
            ).order_by('-bq_last_process_date_time', '-lot_id')

            report_data_pick = []
            for idx, stock_obj in enumerate(pick_queryset, start=1):
                batch = stock_obj.batch_id
                no_of_trays = BrassAuditTrayId.objects.filter(lot_id=stock_obj.lot_id).count()
                data = {
                    'S.No': idx,
                    'Last Updated': stock_obj.bq_last_process_date_time.strftime('%d-%b-%y %I:%M %p') if stock_obj.bq_last_process_date_time else '',
                    'Plating Stk No': batch.plating_stk_no,
                    'Polishing Stk No': batch.polishing_stk_no,
                    'Plating Color': batch.plating_color,
                    'Category': batch.category,
                    'Polish Finish': batch.polish_finish,
                    'Tray Cate-Capacity': f"{batch.tray_type}-{batch.tray_capacity}" if batch.tray_type else '',
                    'Input Source': batch.location.location_name if batch.location else '',
                    'No of Trays': no_of_trays,
                    'Lot Qty': stock_obj.brass_qc_accepted_qty or 0,
                    'Physical Qty': stock_obj.brass_audit_physical_qty or 0,
                    'Accept Qty': stock_obj.brass_audit_accepted_qty or 0,
                    'Reject Qty': stock_obj.brass_rejection_total_qty or 0,
                    'Process Status': 'QC',
                    'Action': 'Delete Disabled   View',
                    'Lot Status': 'Yet to Start',
                    'Current Stage': 'Brass QC',
                    'Remarks': stock_obj.BA_pick_remarks or ''
                }
                report_data_pick.append(data)
            
            df_pick = pd.DataFrame(report_data_pick)
            if len(df_pick) > 0:
                df_pick.to_excel(writer, sheet_name='Pick Table', index=False)
            else:
                # Create empty sheet with headers if no data
                empty_df = pd.DataFrame(columns=['S.No', 'Last Updated', 'Plating Stk No', 'Polishing Stk No', 'Plating Color', 'Category', 'Polish Finish', 'Tray Cate-Capacity', 'Input Source', 'No of Trays', 'Lot Qty', 'Physical Qty', 'Accept Qty', 'Reject Qty', 'Process Status', 'Action', 'Lot Status', 'Current Stage', 'Remarks'])
                empty_df.to_excel(writer, sheet_name='Pick Table', index=False)

            # Completed Table - Similar to BrassAuditCompletedView
            tz = pytz.timezone("Asia/Kolkata")
            now_local = timezone.now().astimezone(tz)
            today = now_local.date()
            yesterday = today - timedelta(days=1)
            from_date = yesterday
            to_date = today
            from_datetime = timezone.make_aware(datetime.combine(from_date, datetime.min.time()))
            to_datetime = timezone.make_aware(datetime.combine(to_date, datetime.max.time()))

            brass_rejection_qty_subquery_comp = Brass_Audit_Rejection_ReasonStore.objects.filter(
                lot_id=OuterRef('lot_id')
            ).values('total_rejection_quantity')[:1]

            completed_queryset = TotalStockModel.objects.select_related(
                'batch_id',
                'batch_id__model_stock_no',
                'batch_id__version',
                'batch_id__location'
            ).filter(
                batch_id__total_batch_quantity__gt=0,
                brass_audit_last_process_date_time__range=(from_datetime, to_datetime)
            ).annotate(
                brass_rejection_qty=brass_rejection_qty_subquery_comp,
            ).filter(
                Q(brass_audit_accptance=True) |
                Q(brass_audit_rejection=True) |
                Q(brass_audit_few_cases_accptance=True, brass_audit_onhold_picking=False) |
                Q(send_brass_audit_to_iqf=True, brass_audit_onhold_picking=False)
            ).order_by('-brass_audit_last_process_date_time')

            report_data_completed = []
            for idx, stock_obj in enumerate(completed_queryset, start=1):
                batch = stock_obj.batch_id
                no_of_trays = BrassAuditTrayId.objects.filter(lot_id=stock_obj.lot_id).count()
                data = {
                    'S.No': idx,
                    'Last Updated': stock_obj.brass_audit_last_process_date_time.strftime('%d-%b-%y %I:%M %p') if stock_obj.brass_audit_last_process_date_time else '',
                    'Plating Stk No': batch.plating_stk_no,
                    'Polishing Stk No': batch.polishing_stk_no,
                    'Plating Color': batch.plating_color,
                    'Category': batch.category,
                    'Polish Finish': batch.polish_finish,
                    'Tray Cate-Capacity': f"{batch.tray_type}-{batch.tray_capacity}" if batch.tray_type else '',
                    'Input Source': batch.location.location_name if batch.location else '',
                    'No of Trays': no_of_trays,
                    'Lot Qty': stock_obj.brass_qc_accepted_qty or 0,
                    'Physical Qty': stock_obj.brass_audit_physical_qty or 0,
                    'Accept Qty': stock_obj.brass_audit_accepted_qty or 0,
                    'Reject Qty': stock_obj.brass_rejection_qty or 0,
                    'Process Status': 'QC',
                    'Action': 'Delete Disabled   View',
                    'Lot Status': 'Completed' if stock_obj.brass_audit_accptance else 'Rejected',
                    'Current Stage': 'Brass Audit',
                    'Remarks': ''
                }
                report_data_completed.append(data)
            
            df_completed = pd.DataFrame(report_data_completed)
            if len(df_completed) > 0:
                df_completed.to_excel(writer, sheet_name='Completed Table', index=False)
            else:
                # Create empty sheet with headers if no data
                empty_df = pd.DataFrame(columns=['S.No', 'Last Updated', 'Plating Stk No', 'Polishing Stk No', 'Plating Color', 'Category', 'Polish Finish', 'Tray Cate-Capacity', 'Input Source', 'No of Trays', 'Lot Qty', 'Physical Qty', 'Accept Qty', 'Reject Qty', 'Process Status', 'Action', 'Lot Status', 'Current Stage', 'Remarks'])
                empty_df.to_excel(writer, sheet_name='Completed Table', index=False)

            # Rejected Table - Aggregated from Brass_Audit_Rejected_TrayScan
            rejected_queryset = Brass_Audit_Rejected_TrayScan.objects.values('lot_id').distinct().order_by('lot_id')

            report_data_rejected = []
            for idx, reject_obj in enumerate(rejected_queryset, start=1):
                # Get batch info from TotalStockModel
                stock_obj = TotalStockModel.objects.filter(lot_id=reject_obj['lot_id']).select_related('batch_id').first()
                if stock_obj:
                    batch = stock_obj.batch_id
                    # Get all rejected trays for this lot
                    rejected_trays = Brass_Audit_Rejected_TrayScan.objects.filter(lot_id=reject_obj['lot_id'])
                    reject_reasons = rejected_trays.values_list('rejection_reason__rejection_reason', flat=True).distinct()
                    # Sum the rejected quantities
                    total_reject_qty = sum(int(tray.rejected_tray_quantity) if tray.rejected_tray_quantity.isdigit() else 0 for tray in rejected_trays)
                    no_of_trays = BrassAuditTrayId.objects.filter(lot_id=reject_obj['lot_id']).count()
                    data = {
                        'S.No': idx,
                        'Last Updated': stock_obj.brass_audit_last_process_date_time.strftime('%d-%b-%y %I:%M %p') if stock_obj.brass_audit_last_process_date_time else '',
                        'Plating Stk No': batch.plating_stk_no,
                        'Polish Stk No': batch.polishing_stk_no,
                        'Plating Color': batch.plating_color,
                        'Polish Finish': batch.polish_finish,
                        'Source - Location': batch.location.location_name if batch.location else '',
                        'Tray Type Capacity': f"{batch.tray_type}-{batch.tray_capacity}" if batch.tray_type else '',
                        'No of Trays': no_of_trays,
                        'Reject Qty': total_reject_qty,
                        'Reject Reason': ', '.join(set(reject_reasons)),
                        'Lot Remark': Brass_Audit_Rejection_ReasonStore.objects.filter(lot_id=reject_obj['lot_id']).first().lot_rejected_comment if Brass_Audit_Rejection_ReasonStore.objects.filter(lot_id=reject_obj['lot_id']).exists() else ''
                    }
                    report_data_rejected.append(data)
            
            df_rejected = pd.DataFrame(report_data_rejected)
            if len(df_rejected) > 0:
                df_rejected.to_excel(writer, sheet_name='Rejected Table', index=False)
            else:
                # Create empty sheet with headers if no data
                empty_df = pd.DataFrame(columns=['S.No', 'Last Updated', 'Plating Stk No', 'Polish Stk No', 'Plating Color', 'Polish Finish', 'Source - Location', 'Tray Type Capacity', 'No of Trays', 'Reject Qty', 'Reject Reason', 'Lot Remark'])
                empty_df.to_excel(writer, sheet_name='Rejected Table', index=False)

        elif module == 'recovery-day-planning':
            # Pick Table
            from Recovery_DP.models import RecoveryTrayId_History
            dp_trays = convert_datetimes(list(RecoveryTrayId_History.objects.all().values()))
            pd.DataFrame(dp_trays).to_excel(writer, sheet_name='Pick Table', index=False)
            
            # Completed Table
            completed_trays = convert_datetimes(list(RecoveryTrayId_History.objects.filter(scanned=True).values()))
            pd.DataFrame(completed_trays).to_excel(writer, sheet_name='Completed Table', index=False)

        elif module == 'recovery-input-screening':
            # Pick Table
            from Recovery_IS.models import RecoveryIPTrayId
            ip_trays = convert_datetimes(list(RecoveryIPTrayId.objects.all().values()))
            pd.DataFrame(ip_trays).to_excel(writer, sheet_name='Pick Table', index=False)
            
            # Accept Table
            from Recovery_IS.models import RecoveryIP_Accepted_TrayScan
            accepted = convert_datetimes(list(RecoveryIP_Accepted_TrayScan.objects.all().values()))
            pd.DataFrame(accepted).to_excel(writer, sheet_name='Accept Table', index=False)
            
            # Reject Table
            from Recovery_IS.models import RecoveryIP_Rejected_TrayScan
            rejected = convert_datetimes(list(RecoveryIP_Rejected_TrayScan.objects.all().values()))
            pd.DataFrame(rejected).to_excel(writer, sheet_name='Reject Table', index=False)
            
            # Completed Table
            from Recovery_IS.models import RecoveryIP_Accepted_TrayID_Store
            completed = convert_datetimes(list(RecoveryIP_Accepted_TrayID_Store.objects.all().values()))
            pd.DataFrame(completed).to_excel(writer, sheet_name='Completed Table', index=False)

        elif module == 'recovery-brass-qc':
            # Pick Table
            from Recovery_Brass_QC.models import BrassTrayId as RecoveryBrassTrayId
            brass_trays = convert_datetimes(list(RecoveryBrassTrayId.objects.all().values()))
            pd.DataFrame(brass_trays).to_excel(writer, sheet_name='Pick Table', index=False)
            
            # Accept Table
            from Recovery_Brass_QC.models import Brass_Qc_Accepted_TrayScan as RecoveryBrass_Qc_Accepted_TrayScan
            accepted = convert_datetimes(list(RecoveryBrass_Qc_Accepted_TrayScan.objects.all().values()))
            pd.DataFrame(accepted).to_excel(writer, sheet_name='Accept Table', index=False)
            
            # Reject Table
            from Recovery_Brass_QC.models import Brass_QC_Rejected_TrayScan as RecoveryBrass_QC_Rejected_TrayScan
            rejected = convert_datetimes(list(RecoveryBrass_QC_Rejected_TrayScan.objects.all().values()))
            pd.DataFrame(rejected).to_excel(writer, sheet_name='Reject Table', index=False)
            
            # Completed Table
            from Recovery_Brass_QC.models import Brass_Qc_Accepted_TrayID_Store as RecoveryBrass_Qc_Accepted_TrayID_Store
            completed = convert_datetimes(list(RecoveryBrass_Qc_Accepted_TrayID_Store.objects.all().values()))
            pd.DataFrame(completed).to_excel(writer, sheet_name='Completed Table', index=False)

        elif module == 'recovery-iqf':
            # Pick Table
            from Recovery_IQF.models import IQFTrayId as RecoveryIQFTrayId
            iqf_trays = convert_datetimes(list(RecoveryIQFTrayId.objects.all().values()))
            pd.DataFrame(iqf_trays).to_excel(writer, sheet_name='Pick Table', index=False)
            
            # Accept Table
            from Recovery_IQF.models import IQF_Accepted_TrayScan as RecoveryIQF_Accepted_TrayScan
            accepted = convert_datetimes(list(RecoveryIQF_Accepted_TrayScan.objects.all().values()))
            pd.DataFrame(accepted).to_excel(writer, sheet_name='Accept Table', index=False)
            
            # Reject Table
            from Recovery_IQF.models import IQF_Rejected_TrayScan as RecoveryIQF_Rejected_TrayScan
            rejected = convert_datetimes(list(RecoveryIQF_Rejected_TrayScan.objects.all().values()))
            pd.DataFrame(rejected).to_excel(writer, sheet_name='Reject Table', index=False)
            
            # Completed Table
            from Recovery_IQF.models import IQF_Accepted_TrayID_Store as RecoveryIQF_Accepted_TrayID_Store
            completed = convert_datetimes(list(RecoveryIQF_Accepted_TrayID_Store.objects.all().values()))
            pd.DataFrame(completed).to_excel(writer, sheet_name='Completed Table', index=False)

        elif module == 'recovery-brass-audit':
            # Pick Table
            from Recovery_BrassAudit.models import BrassAuditTrayId as RecoveryBrassAuditTrayId
            audit_trays = convert_datetimes(list(RecoveryBrassAuditTrayId.objects.all().values()))
            pd.DataFrame(audit_trays).to_excel(writer, sheet_name='Pick Table', index=False)
            
            # Accept Table
            from Recovery_BrassAudit.models import Brass_Audit_Accepted_TrayScan as RecoveryBrass_Audit_Accepted_TrayScan
            accepted = convert_datetimes(list(RecoveryBrass_Audit_Accepted_TrayScan.objects.all().values()))
            pd.DataFrame(accepted).to_excel(writer, sheet_name='Accept Table', index=False)
            
            # Reject Table
            from Recovery_BrassAudit.models import Brass_Audit_Rejected_TrayScan as RecoveryBrass_Audit_Rejected_TrayScan
            rejected = convert_datetimes(list(RecoveryBrass_Audit_Rejected_TrayScan.objects.all().values()))
            pd.DataFrame(rejected).to_excel(writer, sheet_name='Reject Table', index=False)
            
            # Completed Table
            from Recovery_BrassAudit.models import Brass_Audit_Accepted_TrayID_Store as RecoveryBrass_Audit_Accepted_TrayID_Store
            completed = convert_datetimes(list(RecoveryBrass_Audit_Accepted_TrayID_Store.objects.all().values()))
            pd.DataFrame(completed).to_excel(writer, sheet_name='Completed Table', index=False)

        elif module == 'jig-loading':
            _write_jig_loading_report(writer)

        elif module == 'inprocess-inspection':
            _write_inprocess_inspection_report(writer)

        elif module in ('jig-unloading', 'jig-unloading-z1'):
            _write_jig_unloading_report(writer, zone=1)

        elif module == 'jig-unloading-z2':
            _write_jig_unloading_report(writer, zone=2)

        elif module in ('nickel-inspection', 'nickel-inspection-z1'):
            _write_nickel_wiping_report(writer, zone=1)

        elif module == 'nickel-inspection-z2':
            _write_nickel_wiping_report(writer, zone=2)

        elif module in ('nickel-audit', 'nickel-audit-z1'):
            _write_nickel_audit_report(writer, zone=1)

        elif module == 'nickel-audit-z2':
            _write_nickel_audit_report(writer, zone=2)

        elif module == 'spider-spindle-z1':
            _write_spider_spindle_report(writer, zone=1)

        elif module == 'spider-spindle-z2':
            _write_spider_spindle_report(writer, zone=2)

        has_data = _workbook_has_data(writer)
        _autofit_excel_columns(writer)

    if not has_data:
        output.close()
        return HttpResponse('No data found', status=404)

    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={module}_report.xlsx'
    response['Content-Length'] = len(output.getvalue())
    # Clear buffer to prevent buffering
    output.close()
    return response


# ---------------------------------------------------------------------------
# Missing-module report helpers
# ---------------------------------------------------------------------------

def _report_text(value):
    """Normalize display values and prevent spreadsheet formula injection."""
    if value is None:
        return ''
    value = str(value)
    if value.startswith(('=', '+', '-', '@')):
        return f"'{value}"
    return value


def _report_int(value):
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _report_datetime(value):
    if not value:
        return ''
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.strftime('%d-%b-%y %I:%M %p')


def _report_write_sheet(writer, sheet_name, columns, rows):
    """Always write the declared headers, including for an empty report."""
    normalized_rows = []
    for row in rows:
        normalized = {}
        for column in columns:
            value = row.get(column, '')
            normalized[column] = (
                value
                if isinstance(value, (int, float))
                else _report_text(value)
            )
        normalized_rows.append(normalized)
    pd.DataFrame(normalized_rows, columns=columns).to_excel(
        writer,
        sheet_name=sheet_name,
        index=False,
    )


def _report_version(value):
    if not value:
        return ''
    return (
        getattr(value, 'version_internal', None)
        or getattr(value, 'version_name', None)
        or _report_text(value)
    )


def _report_location_names(obj):
    relation = getattr(obj, 'location', None)
    if relation is None:
        return ''
    if hasattr(relation, 'all'):
        return ', '.join(_report_text(item) for item in relation.all())
    return _report_text(relation)


def _report_date_bounds():
    local_today = timezone.localtime(timezone.now()).date()
    from_date = local_today - timedelta(days=1)
    return (
        timezone.make_aware(datetime.combine(from_date, datetime.min.time())),
        timezone.make_aware(datetime.combine(local_today, datetime.max.time())),
    )


def _report_batch_map(records):
    from modelmasterapp.models import ModelMasterCreation

    batch_ids = {record.batch_id for record in records if record.batch_id}
    batches = {
        batch.batch_id: batch
        for batch in ModelMasterCreation.objects.filter(
            batch_id__in=batch_ids
        ).select_related(
            'version',
            'location',
            'model_stock_no',
            'model_stock_no__tray_type',
        )
    }
    missing_batch_ids = batch_ids.difference(batches)
    if missing_batch_ids:
        try:
            from Recovery_DP.models import RecoveryMasterCreation

            for batch in RecoveryMasterCreation.objects.filter(
                batch_id__in=missing_batch_ids
            ).select_related(
                'version',
                'location',
                'model_stock_no',
                'model_stock_no__tray_type',
            ):
                batches[batch.batch_id] = batch
        except (ImportError, AttributeError):
            logger.warning(
                'Recovery batch metadata is unavailable for report export'
            )
    return batches


def _report_stock_map(lot_ids):
    from modelmasterapp.models import TotalStockModel

    stocks = {
        stock.lot_id: stock
        for stock in TotalStockModel.objects.filter(
            lot_id__in=lot_ids
        ).select_related('plating_color', 'batch_id')
    }
    missing_lot_ids = set(lot_ids).difference(stocks)
    if missing_lot_ids:
        try:
            from Recovery_DP.models import RecoveryStockModel

            for stock in RecoveryStockModel.objects.filter(
                lot_id__in=missing_lot_ids
            ).select_related('plating_color', 'batch_id'):
                stocks[stock.lot_id] = stock
        except (ImportError, AttributeError):
            logger.warning(
                'Recovery stock metadata is unavailable for report export'
            )
    return stocks


def _report_model_names(record):
    allocations = record.multi_model_allocation or []
    names = []
    for allocation in allocations:
        if not isinstance(allocation, dict):
            continue
        name = allocation.get('model_name') or allocation.get('model')
        if name:
            names.append(_report_text(name))
    return ', '.join(names) or _report_text(record.plating_stock_num)


def _report_dynamic_capacity(tray_type, fallback=0):
    normalized = _report_text(tray_type).strip().upper()
    if normalized in ('NORMAL', 'NR', 'NB', 'ND', 'NL', 'NW'):
        return 20
    if normalized in ('JUMBO', 'JR', 'JB', 'JD'):
        return 12
    return _report_int(fallback)


def _write_jig_loading_report(writer):
    from collections import Counter
    from django.db.models import Q
    from BrassAudit.models import BrassAuditTrayId
    from Brass_QC.models import BrassTrayId
    from Jig_Loading.models import (
        ExcessLotRecord,
        JigCompleted,
        JigLoadingMaster,
        JigLoadTrayId,
    )
    from modelmasterapp.models import TotalStockModel

    pick_columns = [
        'S.No', 'Last Updated', 'Plating Stk No', 'LOT Qty',
        'No of Trays', 'Process Status', 'Lot Status', 'Current Stage',
        'Polishing Stk No', 'Plating Color', 'Polish Finish', 'Jig Type',
        'In.p Info', 'Version', 'Remarks',
    ]
    completed_columns = [
        'S.No', 'Last Updated', 'Plating Stk No', 'Jig ID', 'LOT Qty',
        'No of Trays', 'Process Status', 'Lot Status', 'Current Stage',
        'Polishing Stk No', 'Plating Color', 'Polish Finish',
        'Tray Type - Capacity', 'Jig Type Capacity', 'Remarks',
    ]

    submitted_records = list(
        JigCompleted.objects.filter(draft_status='submitted').only(
            'lot_id', 'is_multi_model', 'multi_model_allocation'
        )
    )
    submitted_lot_ids = {record.lot_id for record in submitted_records}
    for record in submitted_records:
        for allocation in record.multi_model_allocation or []:
            if isinstance(allocation, dict) and allocation.get('lot_id'):
                submitted_lot_ids.add(allocation['lot_id'])

    draft_by_lot = {
        record.lot_id: record
        for record in JigCompleted.objects.filter(
            draft_status__in=('draft', 'active')
        ).order_by('updated_at')
    }

    stocks = list(
        TotalStockModel.objects.filter(
            Q(brass_audit_accptance=True)
            | Q(
                brass_audit_few_cases_accptance=True,
                brass_audit_onhold_picking=False,
            )
        )
        .exclude(lot_id__in=submitted_lot_ids)
        .select_related(
            'batch_id',
            'batch_id__version',
            'batch_id__model_stock_no',
        )
        .order_by('-brass_audit_last_process_date_time', '-lot_id')
    )
    lot_ids = [stock.lot_id for stock in stocks]

    def tray_counts(model):
        return Counter(
            model.objects.filter(lot_id__in=lot_ids)
            .values_list('lot_id', flat=True)
        )

    jig_counts = tray_counts(JigLoadTrayId)
    audit_counts = tray_counts(BrassAuditTrayId)
    brass_counts = tray_counts(BrassTrayId)
    capacity_map = {
        value.model_stock_no_id: _report_int(value.jig_capacity)
        for value in JigLoadingMaster.objects.filter(
            jig_capacity__isnull=False
        )
    }

    pick_rows = []
    for stock in stocks:
        batch = stock.batch_id
        draft = draft_by_lot.get(stock.lot_id)
        lot_status = (
            'Draft' if draft
            else 'On Hold' if stock.jig_hold_lot
            else 'Yet to Start'
        )
        current_stage = (
            'Jig Loading'
            if draft or stock.jig_hold_lot else 'Brass Audit'
        )
        model_id = getattr(batch, 'model_stock_no_id', None)
        pick_rows.append({
            'S.No': len(pick_rows) + 1,
            'Last Updated': _report_datetime(
                stock.brass_audit_last_process_date_time
            ),
            'Plating Stk No': batch.plating_stk_no,
            'LOT Qty': _report_int(
                stock.brass_audit_accepted_qty
                or stock.brass_audit_physical_qty
                or stock.total_stock
            ),
            'No of Trays': (
                jig_counts.get(stock.lot_id)
                or audit_counts.get(stock.lot_id)
                or brass_counts.get(stock.lot_id)
                or 0
            ),
            'Process Status': 'L',
            'Lot Status': lot_status,
            'Current Stage': current_stage,
            'Polishing Stk No': batch.polishing_stk_no,
            'Plating Color': batch.plating_color,
            'Polish Finish': batch.polish_finish,
            'Jig Type': f"{capacity_map.get(model_id, 0)} Jig",
            'In.p Info': getattr(draft, 'pick_remarks', '') if draft else '',
            'Version': _report_version(batch.version),
            'Remarks': stock.jig_pick_remarks or stock.jig_holding_reason,
        })

    excess_records = list(
        ExcessLotRecord.objects.exclude(
            new_lot_id__in=submitted_lot_ids
        ).select_related(
            'jig_loading_record'
        ).prefetch_related(
            'excess_trays'
        ).order_by('-created_at')
    )
    parent_ids = {record.parent_lot_id for record in excess_records}
    parent_stocks = {
        stock.lot_id: stock
        for stock in TotalStockModel.objects.filter(
            lot_id__in=parent_ids
        ).select_related('batch_id', 'batch_id__version')
    }
    for record in excess_records:
        stock = parent_stocks.get(record.parent_lot_id)
        batch = getattr(stock, 'batch_id', None)
        if not batch:
            continue
        pick_rows.append({
            'S.No': len(pick_rows) + 1,
            'Last Updated': _report_datetime(record.created_at),
            'Plating Stk No': batch.plating_stk_no,
            'LOT Qty': _report_int(record.lot_qty),
            'No of Trays': len(record.excess_trays.all()),
            'Process Status': 'L',
            'Lot Status': 'Yet to Start',
            'Current Stage': 'Jig Loading',
            'Polishing Stk No': batch.polishing_stk_no,
            'Plating Color': batch.plating_color,
            'Polish Finish': batch.polish_finish,
            'Jig Type': (
                f"{_report_int(record.jig_loading_record.jig_capacity)} Jig"
            ),
            'In.p Info': 'Excess Lot',
            'Version': _report_version(batch.version),
            'Remarks': '',
        })

    completed_records = list(
        JigCompleted.objects.filter(draft_status='submitted')
        .select_related('user')
        .order_by('-updated_at')
    )
    batches = _report_batch_map(completed_records)
    completed_rows = []
    for record in completed_records:
        batch = batches.get(record.batch_id)
        completed_rows.append({
            'S.No': len(completed_rows) + 1,
            'Last Updated': _report_datetime(record.updated_at),
            'Plating Stk No': (
                _report_model_names(record)
                or getattr(batch, 'plating_stk_no', '')
            ),
            'Jig ID': record.jig_id,
            'LOT Qty': _report_int(
                record.delink_tray_qty or record.loaded_cases_qty
            ),
            'No of Trays': _report_int(record.delink_tray_count),
            'Process Status': 'L',
            'Lot Status': 'Released',
            'Current Stage': 'Jig Loading',
            'Polishing Stk No': getattr(batch, 'polishing_stk_no', ''),
            'Plating Color': getattr(batch, 'plating_color', ''),
            'Polish Finish': getattr(batch, 'polish_finish', ''),
            'Tray Type - Capacity': (
                f"{_report_text(record.tray_type)} - "
                f"{_report_int(record.tray_capacity)}"
                if record.tray_type else ''
            ),
            'Jig Type Capacity': f"Jig-{_report_int(record.jig_capacity)}",
            'Remarks': record.remarks,
        })

    _report_write_sheet(writer, 'Pick Table', pick_columns, pick_rows)
    _report_write_sheet(
        writer, 'Completed Table', completed_columns, completed_rows
    )


def _inprocess_rows(records, completed=False):
    batches = _report_batch_map(records)
    stock_map = _report_stock_map(
        [record.lot_id for record in records]
    )
    rows = []
    for record in records:
        batch = batches.get(record.batch_id)
        stock = stock_map.get(record.lot_id)
        hold = bool(getattr(stock, 'inprocess_hold_lot', False))
        release = bool(getattr(stock, 'inprocess_release_lot', False))
        lot_status = (
            'On Hold' if hold
            else 'Released' if release or completed
            else 'Yet to Start'
        )
        model_names = _report_model_names(record)
        lot_qty = _report_int(
            record.original_lot_qty
            or record.delink_tray_qty
            or record.loaded_cases_qty
        )
        common = {
            'S.No': len(rows) + 1,
            'JIG ID': record.jig_id,
            'Date & Time': _report_datetime(record.updated_at),
            'Model Presents': model_names,
            'Plating Stk No': getattr(batch, 'plating_stk_no', ''),
            'Plating Color': getattr(batch, 'plating_color', ''),
            'Polish Finish': getattr(batch, 'polish_finish', ''),
            'Version': _report_version(getattr(batch, 'version', None)),
            'Remarks': record.pick_remarks or record.remarks,
        }
        if completed:
            common.update({
                'Polishing Stk No': getattr(
                    batch, 'polishing_stk_no', ''
                ),
                'Source- Location': _report_text(
                    getattr(batch, 'location', '')
                ),
                'Tray Type - Capacity': (
                    f"{_report_text(record.tray_type)} - "
                    f"{_report_int(record.tray_capacity)}"
                    if record.tray_type else ''
                ),
                'Jig Type - Capacity': (
                    f"Jig - {_report_int(record.jig_capacity)}"
                ),
                'Bath Type': (
                    record.nickel_bath_type
                    or getattr(batch, 'ep_bath_type', '')
                ),
                'Jig Lot Qty': lot_qty,
                'Bath No': getattr(record.bath_numbers, 'bath_number', ''),
                'IP Info': record.jig_position,
                'Process Status': 'Completed',
                'Batch Status': 'Completed',
                'Current Stage': 'Inprocess Inspection',
            })
        else:
            common.update({
                'Nickel Bath No': getattr(
                    record.bath_numbers, 'bath_number', ''
                ),
                'Bath Type': (
                    record.nickel_bath_type
                    or getattr(batch, 'ep_bath_type', '')
                ),
                'Process Status': 'IP',
                'Jig Cate-Capacity': (
                    f"Jig - {_report_int(record.jig_capacity)}"
                ),
                'Lot Qty': lot_qty,
                'Lot Status': lot_status,
                'Current Stage': 'Inprocess Inspection',
                'Tray Cate-Capacity': (
                    f"{_report_text(record.tray_type)} - "
                    f"{_report_int(record.tray_capacity)}"
                    if record.tray_type else ''
                ),
                'In.P Info': record.jig_position or '',
            })
        rows.append(common)
    return rows


def _write_inprocess_inspection_report(writer):
    from Jig_Loading.models import JigCompleted

    pick_columns = [
        'S.No', 'JIG ID', 'Date & Time', 'Model Presents',
        'Nickel Bath No', 'Bath Type', 'Process Status',
        'Jig Cate-Capacity', 'Lot Qty', 'Lot Status', 'Current Stage',
        'Plating Stk No', 'Plating Color', 'Polish Finish',
        'Tray Cate-Capacity', 'In.P Info', 'Version', 'Remarks',
    ]
    completed_columns = [
        'S.No', 'JIG ID', 'Date & Time', 'Model Presents',
        'Plating Stk No', 'Polishing Stk No', 'Plating Color',
        'Polish Finish', 'Version', 'Source- Location',
        'Tray Type - Capacity', 'Jig Type - Capacity', 'Bath Type',
        'Jig Lot Qty', 'Bath No', 'IP Info', 'Process Status',
        'Batch Status', 'Current Stage', 'Remarks',
    ]
    pick_records = list(
        JigCompleted.objects.filter(
            jig_position__isnull=True,
            draft_status='submitted',
        ).select_related('bath_numbers').order_by('-updated_at')
    )
    from_datetime, to_datetime = _report_date_bounds()
    completed_records = list(
        JigCompleted.objects.filter(
            jig_position__isnull=False,
            updated_at__range=(from_datetime, to_datetime),
        ).select_related('bath_numbers').order_by('-updated_at')
    )
    _report_write_sheet(
        writer, 'IP Main', pick_columns, _inprocess_rows(pick_records)
    )
    _report_write_sheet(
        writer,
        'IP Completed',
        completed_columns,
        _inprocess_rows(completed_records, completed=True),
    )


def _normalize_unload_lot_id(value):
    value = _report_text(value).strip().lstrip('-')
    if ':' in value:
        value = value.rsplit(':', 1)[-1].strip()
    if value.startswith('JLOT-') and '-' in value[5:]:
        value = value.rsplit('-', 1)[-1]
    return value


def _represented_jig_lots(record):
    represented = set()
    draft_data = record.draft_data if isinstance(record.draft_data, dict) else {}
    quantities = draft_data.get('lot_id_quantities', {})
    if isinstance(quantities, dict):
        represented.update(
            _normalize_unload_lot_id(value) for value in quantities
        )
    for allocation in record.multi_model_allocation or []:
        if isinstance(allocation, dict) and allocation.get('lot_id'):
            represented.add(_normalize_unload_lot_id(allocation['lot_id']))
    if record.lot_id:
        represented.add(_normalize_unload_lot_id(record.lot_id))
    return {value for value in represented if value}


def _jig_unload_color(record, batch, stock):
    draft_data = record.draft_data if isinstance(record.draft_data, dict) else {}
    color = (
        draft_data.get('plating_color')
        or getattr(batch, 'plating_color', '')
        or _report_text(getattr(stock, 'plating_color', ''))
    )
    color = _report_text(color).strip()
    return color[3:] if color.upper().startswith('IP-') else color


def _write_jig_unloading_report(writer, zone):
    from collections import defaultdict
    from Jig_Loading.models import JigCompleted
    from Jig_Unloading.models import JigUnloadAfterTable
    from modelmasterapp.models import TotalStockModel

    pick_columns = [
        'S.No', 'JIG ID', 'Last Updated', 'Lot Qty', 'Model Presents',
        'Bath No', 'Process Status', 'Lot Status', 'Current Stage',
        'Polish Finish', 'Remarks',
    ]
    completed_columns = [
        'S.No', 'JIG ID', 'Last Updated', 'Model Presents', 'Lot Qty',
        'No of Trays', 'Process Status', 'Batch Status',
        'Current Location', 'Plating Stk No', 'Polishing Stk No',
        'Plating Color', 'Polish Finish', 'Tray Cate - Capacity',
        'Version', 'Remarks',
    ]

    process_modules = (
        ['Inprocess Inspection']
        if zone == 1 else ['Inprocess Inspection', 'Jig Unloading']
    )
    candidates = list(
        JigCompleted.objects.filter(
            last_process_module__in=process_modules
        ).select_related('bath_numbers').order_by('-IP_loaded_date_time')
    )
    batches = _report_batch_map(candidates)
    stocks = _report_stock_map(
        [record.lot_id for record in candidates]
    )

    unloaded_by_jig = defaultdict(set)
    for unload in JigUnloadAfterTable.objects.only(
        'jig_qr_id', 'combine_lot_ids'
    ):
        for combined in unload.combine_lot_ids or []:
            unloaded_by_jig[unload.jig_qr_id].add(
                _normalize_unload_lot_id(combined)
            )

    pick_rows = []
    for record in candidates:
        batch = batches.get(record.batch_id)
        stock = stocks.get(record.lot_id)
        color = _jig_unload_color(record, batch, stock)
        is_zone_one = color.upper() == 'IPS'
        if is_zone_one != (zone == 1):
            continue
        represented = _represented_jig_lots(record)
        if represented and represented.issubset(
            unloaded_by_jig.get(record.jig_id, set())
        ):
            continue
        hold = (record.hold_status or '').lower() == 'hold'
        pick_rows.append({
            'S.No': len(pick_rows) + 1,
            'JIG ID': record.jig_id,
            'Last Updated': _report_datetime(
                record.IP_loaded_date_time or record.updated_at
            ),
            'Lot Qty': _report_int(
                record.updated_lot_qty
                or record.original_lot_qty
                or record.loaded_cases_qty
                or record.delink_tray_qty
            ),
            'Model Presents': _report_model_names(record),
            'Bath No': getattr(record.bath_numbers, 'bath_number', ''),
            'Process Status': 'JU',
            'Lot Status': 'On Hold' if hold else 'Yet to Start',
            'Current Stage': record.last_process_module,
            'Polish Finish': getattr(batch, 'polish_finish', ''),
            'Remarks': record.unloading_remarks or record.pick_remarks,
        })

    from_datetime, to_datetime = _report_date_bounds()
    completed_records = list(
        JigUnloadAfterTable.objects.filter(
            Un_loaded_date_time__range=(from_datetime, to_datetime)
        ).select_related(
            'version', 'plating_color', 'polish_finish'
        ).prefetch_related('location').order_by('-Un_loaded_date_time')
    )

    all_combined_lots = {
        _normalize_unload_lot_id(value)
        for record in completed_records
        for value in (record.combine_lot_ids or [])
    }
    combined_stocks = _report_stock_map(all_combined_lots)
    jig_ids = {record.jig_qr_id for record in completed_records}
    remarks_by_jig = {
        record.jig_id: record.unloading_remarks
        for record in JigCompleted.objects.filter(
            jig_id__in=jig_ids
        ).order_by('updated_at')
    }

    completed_rows = []
    for record in completed_records:
        color = _report_text(record.plating_color)
        if not color:
            for combined in record.combine_lot_ids or []:
                stock = combined_stocks.get(
                    _normalize_unload_lot_id(combined)
                )
                if stock and stock.plating_color:
                    color = _report_text(stock.plating_color)
                    break
                if stock and stock.batch_id:
                    color = _report_text(stock.batch_id.plating_color)
                    if color:
                        break
        is_zone_one = color.upper().replace('IP-', '') == 'IPS'
        if is_zone_one != (zone == 1):
            continue

        plating_values = record.plating_stk_no_list or []
        if not plating_values and record.plating_stk_no:
            plating_values = [record.plating_stk_no]
        model_presents = ', '.join(
            _report_text(value) for value in plating_values
        )
        capacity = _report_dynamic_capacity(
            record.tray_type, record.tray_capacity
        )
        total_qty = _report_int(record.total_case_qty)
        no_of_trays = (
            (total_qty + capacity - 1) // capacity if capacity else 0
        )
        completed_rows.append({
            'S.No': len(completed_rows) + 1,
            'JIG ID': record.jig_qr_id,
            'Last Updated': _report_datetime(record.Un_loaded_date_time),
            'Model Presents': model_presents,
            'Lot Qty': total_qty,
            'No of Trays': no_of_trays,
            'Process Status': 'JU',
            'Batch Status': 'Completed',
            'Current Location': _report_location_names(record),
            'Plating Stk No': ', '.join(
                _report_text(value) for value in plating_values
            ),
            'Polishing Stk No': ', '.join(
                _report_text(value) for value in (
                    record.polish_stk_no_list or [record.polish_stk_no]
                ) if value
            ),
            'Plating Color': color,
            'Polish Finish': record.polish_finish,
            'Tray Cate - Capacity': (
                f"{_report_text(record.tray_type)} - {capacity}"
                if record.tray_type else ''
            ),
            'Version': ', '.join(
                _report_text(value) for value in (
                    record.version_list or [_report_version(record.version)]
                ) if value
            ),
            'Remarks': remarks_by_jig.get(record.jig_qr_id, ''),
        })

    _report_write_sheet(writer, 'JUL Main Table', pick_columns, pick_rows)
    _report_write_sheet(
        writer, 'JUL Completed', completed_columns, completed_rows
    )


def _zone_color_ids(zone):
    from modelmasterapp.models import Plating_Color

    field = (
        'jig_unload_zone_1' if zone == 1 else 'jig_unload_zone_2'
    )
    return list(
        Plating_Color.objects.filter(**{field: True})
        .values_list('id', flat=True)
    )


def _rejection_total_map(model, lot_ids):
    totals = {}
    for lot_id, quantity in model.objects.filter(
        lot_id__in=lot_ids
    ).values_list('lot_id', 'total_rejection_quantity'):
        totals.setdefault(lot_id, _report_int(quantity))
    return totals


def _nickel_lot_status(record, prefix):
    if getattr(record, f'{prefix}_hold_lot', False):
        return 'On Hold'
    if getattr(record, f'{prefix}_qc_rejection', False):
        return 'Rejected'
    if getattr(record, f'{prefix}_qc_few_cases_accptance', False):
        return 'Partially Accepted'
    if getattr(record, f'{prefix}_qc_accptance', False):
        return 'Accepted'
    if getattr(record, f'{prefix}_release_lot', False):
        return 'Released'
    return 'Yet to Start'


def _nickel_process_status(record, prefix):
    if getattr(record, f'{prefix}_draft', False):
        return 'Draft'
    if getattr(record, f'{prefix}_qc_rejection', False):
        return 'Rejected'
    if getattr(record, f'{prefix}_qc_few_cases_accptance', False):
        return 'Partially Accepted'
    if getattr(record, f'{prefix}_qc_accptance', False):
        return 'Accepted'
    return 'Yet to Start'


def _write_nickel_wiping_report(writer, zone):
    from django.db.models import Q
    from Jig_Unloading.models import JigUnloadAfterTable
    from Nickel_Inspection.models import (
        NickelQC_PartialAcceptLot,
        Nickel_QC_Rejection_ReasonStore,
    )

    table_columns = [
        'S.No', 'Last Updated', 'Plating Stk No', 'No of Trays',
        'Lot Qty', 'Accept Qty', 'Reject Qty', 'Lot Status',
        'Current Stage', 'Process Status', 'Polishing Stk No',
        'Plating Color', 'Polish Finish', 'Input Source', 'Remarks',
        'Version',
    ]
    reject_columns = [
        'S.No', 'Last Updated', 'Plating Stk No',
        'Tray Cate- Capacity', 'No of Trays', 'Lot Qty', 'Reject Qty',
        'Lot Status', 'Current Stage', 'Process Status',
        'Polishing Stk No', 'Plating Color', 'Polish Finish',
        'Input Source', 'Remarks', 'Version',
    ]
    base = JigUnloadAfterTable.objects.filter(
        total_case_qty__gt=0,
        plating_color_id__in=_zone_color_ids(zone),
    ).select_related(
        'version', 'plating_color', 'polish_finish'
    ).prefetch_related('location')

    pick_records = list(
        base.filter(
            Q(nq_qc_accptance=False) | Q(nq_qc_accptance__isnull=True),
            Q(nq_qc_rejection=False) | Q(nq_qc_rejection__isnull=True),
        ).exclude(
            nq_qc_few_cases_accptance=True,
            nq_onhold_picking=False,
        ).order_by('-created_at', '-lot_id')
    )
    completed_query = base.filter(
        Q(nq_qc_accptance=True)
        | Q(nq_qc_rejection=True)
        | Q(nq_qc_few_cases_accptance=True, nq_onhold_picking=False)
    ).exclude(
        lot_id__in=NickelQC_PartialAcceptLot.objects.values_list(
            'new_lot_id', flat=True
        )
    )
    if zone == 2:
        from_datetime, to_datetime = _report_date_bounds()
        completed_query = completed_query.filter(
            nq_last_process_date_time__range=(from_datetime, to_datetime)
        )
    completed_records = list(
        completed_query.order_by('-nq_last_process_date_time', '-lot_id')
    )
    reject_records = list(
        base.filter(
            Q(nq_qc_rejection=True)
            | Q(nq_qc_few_cases_accptance=True)
        ).order_by('-nq_last_process_date_time', '-lot_id')
    )
    all_records = pick_records + completed_records + reject_records
    rejection_totals = _rejection_total_map(
        Nickel_QC_Rejection_ReasonStore,
        [record.lot_id for record in all_records],
    )

    def rows(records, reject=False):
        output = []
        for record in records:
            quantity = (
                _report_int(record.nq_physical_qty)
                if reject else _report_int(record.total_case_qty)
            )
            capacity = _report_dynamic_capacity(
                record.tray_type, record.tray_capacity
            )
            no_of_trays = (
                (quantity + capacity - 1) // capacity
                if quantity and capacity else 0
            )
            row = {
                'S.No': len(output) + 1,
                'Last Updated': _report_datetime(
                    record.nq_last_process_date_time or record.created_at
                ),
                'Plating Stk No': record.plating_stk_no,
                'No of Trays': no_of_trays,
                'Lot Qty': quantity,
                'Accept Qty': _report_int(record.nq_qc_accepted_qty),
                'Reject Qty': rejection_totals.get(record.lot_id, 0),
                'Lot Status': _nickel_lot_status(record, 'nq'),
                'Current Stage': (
                    record.current_stage or record.last_process_module
                ),
                'Process Status': _nickel_process_status(record, 'nq'),
                'Polishing Stk No': record.polish_stk_no,
                'Plating Color': record.plating_color,
                'Polish Finish': record.polish_finish,
                'Input Source': _report_location_names(record),
                'Remarks': record.nq_pick_remarks,
                'Version': _report_version(record.version),
            }
            if reject:
                row['Tray Cate- Capacity'] = (
                    f"{_report_text(record.tray_type)} - {capacity}"
                    if record.tray_type else ''
                )
            output.append(row)
        return output

    _report_write_sheet(
        writer, 'Main Table', table_columns, rows(pick_records)
    )
    _report_write_sheet(
        writer, 'Completed Table', table_columns, rows(completed_records)
    )
    _report_write_sheet(
        writer, 'Rejected Table', reject_columns, rows(reject_records, True)
    )


def _audit_source_key(record):
    combined = tuple(
        sorted(
            _normalize_unload_lot_id(value)
            for value in (record.combine_lot_ids or [])
            if value
        )
    )
    return combined or (record.lot_id,)


def _unique_audit_records(records, completed=False):
    selected = {}
    for record in records:
        key = _audit_source_key(record)
        current = selected.get(key)
        if current is None:
            selected[key] = record
            continue
        if not completed:
            continue
        record_priority = (
            0 if record.na_qc_rejection
            else 1 if (
                record.na_qc_few_cases_accptance
                and not record.na_onhold_picking
            )
            else 2
        )
        current_priority = (
            0 if current.na_qc_rejection
            else 1 if (
                current.na_qc_few_cases_accptance
                and not current.na_onhold_picking
            )
            else 2
        )
        if record_priority < current_priority:
            selected[key] = record
    return list(selected.values())


def _write_nickel_audit_report(writer, zone):
    from django.db.models import Q
    from Jig_Unloading.models import JigUnloadAfterTable
    from Nickel_Audit.models import (
        NickelAudit_PartialAcceptLot,
        NickelAudit_Submission,
        Nickel_Audit_Rejection_ReasonStore,
    )
    from Nickel_Inspection.models import NickelQC_PartialAcceptLot

    pick_columns = [
        'S.No', 'Last Updated', 'Plating Stk No', 'No of Trays',
        'Lot Qty', 'Accept Qty', 'Reject Qty', 'Lot Status',
        'Current Stage', 'Process Status', 'Polishing Stk No',
        'Plating Color', 'Polish Finish', 'Input Source', 'Remarks',
        'Version',
    ]
    completed_columns = [
        'S.No', 'Last Updated', 'Plating Stk No', 'Polishing Stk No',
        'Plating Color', 'Category', 'Polish Finish',
        'Tray Cate- Capacity', 'Input Source', 'No of Trays', 'Lot Qty',
        'Physical Qty', 'Accept Qty', 'Reject Qty', 'Process Status',
        'Lot Status', 'Current Stage', 'Remarks',
    ]
    base = JigUnloadAfterTable.objects.filter(
        total_case_qty__gt=0,
        plating_color_id__in=_zone_color_ids(zone),
    ).select_related(
        'version', 'plating_color', 'polish_finish'
    ).prefetch_related('location')

    nq_partial_children = set(
        NickelQC_PartialAcceptLot.objects.values_list(
            'new_lot_id', flat=True
        )
    )
    pick_candidates = list(
        base.filter(
            (
                (
                    Q(na_qc_accptance=False)
                    | Q(na_qc_accptance__isnull=True)
                )
                & (
                    Q(na_qc_rejection=False)
                    | Q(na_qc_rejection__isnull=True)
                )
                & ~Q(
                    na_qc_few_cases_accptance=True,
                    na_onhold_picking=False,
                )
                & (
                    Q(nq_qc_accptance=True)
                    | Q(lot_id__in=nq_partial_children)
                    | Q(
                        nq_qc_few_cases_accptance=True,
                        nq_onhold_picking=False,
                    )
                )
            )
            | Q(na_qc_rejection=True, na_onhold_picking=True)
        ).distinct().order_by('-nq_last_process_date_time', '-lot_id')
    )
    submitted_lots = set(
        NickelAudit_Submission.objects.values_list('lot_id', flat=True)
    )
    visible_lot_ids = {record.lot_id for record in pick_candidates}
    partial_links = list(
        NickelQC_PartialAcceptLot.objects.filter(
            Q(parent_lot_id__in=visible_lot_ids)
            | Q(new_lot_id__in=visible_lot_ids)
        ).values('parent_lot_id', 'new_lot_id')
    )
    visible_child_ids = {
        link['new_lot_id'] for link in partial_links
        if link['new_lot_id'] in visible_lot_ids
    }
    parent_ids_with_visible_child = {
        link['parent_lot_id'] for link in partial_links
        if link['new_lot_id'] in visible_child_ids
    }
    completed_source_lots = set()
    completed_source_query = base.filter(
        Q(na_qc_accptance=True)
        | Q(na_qc_rejection=True)
        | Q(na_qc_few_cases_accptance=True, na_onhold_picking=False)
    ).values_list('lot_id', 'combine_lot_ids')
    for lot_id, combine_lot_ids in completed_source_query:
        source_lots = tuple(
            _normalize_unload_lot_id(value)
            for value in (combine_lot_ids or [])
            if value
        ) or (lot_id,)
        completed_source_lots.update(source_lots)

    pick_records = []
    seen_pick_sources = set()
    for record in pick_candidates:
        if record.lot_id in submitted_lots:
            continue
        if (
            record.lot_id in parent_ids_with_visible_child
            and record.lot_id not in visible_child_ids
        ):
            continue
        source_lots = set(_audit_source_key(record))
        if source_lots.intersection(completed_source_lots):
            continue
        if source_lots.intersection(seen_pick_sources):
            continue
        seen_pick_sources.update(source_lots)
        pick_records.append(record)

    from_datetime, to_datetime = _report_date_bounds()
    completed_query = base.filter(
        Q(na_qc_accptance=True)
        | Q(na_qc_rejection=True)
        | Q(na_qc_few_cases_accptance=True, na_onhold_picking=False),
        na_last_process_date_time__range=(from_datetime, to_datetime),
    ).exclude(
        lot_id__in=NickelAudit_PartialAcceptLot.objects.values_list(
            'new_lot_id', flat=True
        )
    ).order_by('-na_last_process_date_time', '-lot_id')
    completed_records = _unique_audit_records(
        list(completed_query), completed=True
    )

    all_records = pick_records + completed_records
    rejection_totals = _rejection_total_map(
        Nickel_Audit_Rejection_ReasonStore,
        [record.lot_id for record in all_records],
    )
    latest_submissions = {}
    for submission in NickelAudit_Submission.objects.filter(
        lot_id__in=[record.lot_id for record in completed_records]
    ).order_by('-created_at'):
        latest_submissions.setdefault(submission.lot_id, submission)

    pick_rows = []
    for record in pick_records:
        reject_qty = rejection_totals.get(record.lot_id, 0)
        lot_qty = max(
            _report_int(record.nq_qc_accepted_qty) - reject_qty, 0
        )
        capacity = _report_dynamic_capacity(
            record.tray_type, record.tray_capacity
        )
        pick_rows.append({
            'S.No': len(pick_rows) + 1,
            'Last Updated': _report_datetime(
                record.nq_last_process_date_time or record.created_at
            ),
            'Plating Stk No': record.plating_stk_no,
            'No of Trays': (
                (lot_qty + capacity - 1) // capacity
                if lot_qty and capacity else 0
            ),
            'Lot Qty': lot_qty,
            'Accept Qty': _report_int(record.na_qc_accepted_qty),
            'Reject Qty': reject_qty,
            'Lot Status': _nickel_lot_status(record, 'na'),
            'Current Stage': (
                record.current_stage or record.last_process_module
            ),
            'Process Status': _nickel_process_status(record, 'na'),
            'Polishing Stk No': record.polish_stk_no,
            'Plating Color': record.plating_color,
            'Polish Finish': record.polish_finish,
            'Input Source': _report_location_names(record),
            'Remarks': record.na_pick_remarks,
            'Version': _report_version(record.version),
        })

    completed_rows = []
    for record in completed_records:
        submission = latest_submissions.get(record.lot_id)
        accepted_qty = _report_int(
            getattr(submission, 'accepted_qty', None)
            if submission else record.na_qc_accepted_qty
        )
        rejected_qty = _report_int(
            getattr(submission, 'rejected_qty', None)
            if submission else rejection_totals.get(record.lot_id, 0)
        )
        lot_qty = accepted_qty
        capacity = _report_dynamic_capacity(
            record.tray_type, record.tray_capacity
        )
        completed_rows.append({
            'S.No': len(completed_rows) + 1,
            'Last Updated': _report_datetime(
                record.na_last_process_date_time
            ),
            'Plating Stk No': record.plating_stk_no,
            'Polishing Stk No': record.polish_stk_no,
            'Plating Color': record.plating_color,
            'Category': record.category,
            'Polish Finish': record.polish_finish,
            'Tray Cate- Capacity': (
                f"{_report_text(record.tray_type)} - {capacity}"
                if record.tray_type else ''
            ),
            'Input Source': _report_location_names(record),
            'No of Trays': (
                (lot_qty + capacity - 1) // capacity
                if lot_qty and capacity else 0
            ),
            'Lot Qty': lot_qty,
            'Physical Qty': _report_int(record.na_physical_qty),
            'Accept Qty': accepted_qty,
            'Reject Qty': rejected_qty,
            'Process Status': _nickel_process_status(record, 'na'),
            'Lot Status': _nickel_lot_status(record, 'na'),
            'Current Stage': (
                record.current_stage or record.last_process_module
            ),
            'Remarks': record.na_pick_remarks,
        })

    _report_write_sheet(writer, 'NA Pick Table', pick_columns, pick_rows)
    _report_write_sheet(
        writer, 'NA Completed', completed_columns, completed_rows
    )


def _write_spider_spindle_report(writer, zone):
    from collections import defaultdict
    from Jig_Unloading.models import JigUnloadAfterTable
    from SpiderSpindle_Z1.models import SpiderSpindleZ1TrayId
    from SpiderSpindle_Z2.models import SpiderSpindleZ2TrayId

    pick_columns = [
        'S.No', 'Date & Time', 'Plating Stk No', 'Polishing Stk No',
        'Plating Color', 'Category', 'Polish Finish', 'Version',
        'Tray Type', 'Source', 'Input Qty', 'Remarks',
    ]
    completed_columns = [
        'S.No', 'Lot ID', 'Completed At', 'Plating Stk No',
        'Polishing Stk No', 'Plating Color', 'Category', 'Polish Finish',
        'Version', 'Tray Type', 'Source', 'Input Qty', 'Tray ID',
        'Remarks',
    ]
    completed_field = f'ss_z{zone}_completed'
    completed_at_field = f'ss_z{zone}_completed_at'
    tray_id_field = f'ss_z{zone}_tray_id'

    base = JigUnloadAfterTable.objects.filter(
        total_case_qty__gt=0,
        plating_color_id__in=_zone_color_ids(zone),
        na_qc_accptance=True,
    ).select_related(
        'version', 'plating_color', 'polish_finish'
    ).prefetch_related('location')

    pick_records = list(
        base.filter(**{completed_field: False})
        .order_by('-created_at', '-lot_id')
    )
    from_datetime, to_datetime = _report_date_bounds()
    completed_records = list(
        base.filter(
            **{
                completed_field: True,
                f'{completed_at_field}__range': (
                    from_datetime, to_datetime
                ),
            }
        ).order_by(f'-{completed_at_field}', '-lot_id')
    )
    tray_model = (
        SpiderSpindleZ1TrayId if zone == 1 else SpiderSpindleZ2TrayId
    )
    tray_map = defaultdict(list)
    for lot_id, tray_id in tray_model.objects.filter(
        lot_id__in=[record.lot_id for record in completed_records]
    ).order_by('linked_at', 'id').values_list('lot_id', 'tray_id'):
        tray_map[lot_id].append(_report_text(tray_id))

    def common(record):
        return {
            'Plating Stk No': record.plating_stk_no,
            'Polishing Stk No': record.polish_stk_no,
            'Plating Color': record.plating_color,
            'Category': record.category,
            'Polish Finish': record.polish_finish,
            'Version': _report_version(record.version),
            'Tray Type': record.tray_type,
            'Source': _report_location_names(record),
            'Input Qty': _report_int(record.total_case_qty),
            'Remarks': record.spider_pick_remarks,
        }

    pick_rows = []
    for record in pick_records:
        row = {
            'S.No': len(pick_rows) + 1,
            'Date & Time': _report_datetime(
                record.na_last_process_date_time or record.created_at
            ),
        }
        row.update(common(record))
        pick_rows.append(row)

    completed_rows = []
    for record in completed_records:
        tray_ids = tray_map.get(record.lot_id)
        if not tray_ids:
            tray_ids = [
                value.strip()
                for value in _report_text(
                    getattr(record, tray_id_field)
                ).split(',')
                if value.strip()
            ]
        row = {
            'S.No': len(completed_rows) + 1,
            'Lot ID': record.lot_id,
            'Completed At': _report_datetime(
                getattr(record, completed_at_field)
            ),
            'Tray ID': ', '.join(tray_ids),
        }
        row.update(common(record))
        completed_rows.append(row)

    _report_write_sheet(writer, 'Pick Table', pick_columns, pick_rows)
    _report_write_sheet(
        writer, 'Completed Table', completed_columns, completed_rows
    )


# ---------------------------------------------------------------------------
# Consolidated Report (Preview + Download share the same selector)
# ---------------------------------------------------------------------------

CONSOLIDATED_COLUMNS = [
    'S.No', 'Plating Stock No.', 'Lot Qty',
    'Day Planning', 'Input Screening', 'Brass QC', 'IQF', 'Brass Audit',
    'Jig Loading', 'IP Inspection',
    'Jig Unloading Z1', 'Jig Unloading Z2',
    'Nickel Wiping Z1', 'Nickel Wiping Z2',
    'Nickel Audit Z1', 'Nickel Audit Z2',
    'Spider Spindle Z1', 'Spider Spindle Z2',
    'Remarks',
]


def _parse_report_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(value.strip(), '%Y-%m-%d').date()
    except (ValueError, AttributeError):
        return None


def _consolidated_rows_from_request(request):
    from .selectors import get_consolidated_report_rows

    date_from = _parse_report_date(request.GET.get('date_from'))
    date_to = _parse_report_date(request.GET.get('date_to'))
    plating_stock_no = (request.GET.get('plating_stk_no') or '').strip()
    module = (request.GET.get('module') or '').strip()
    if date_from and date_to and date_from > date_to:
        date_from, date_to = date_to, date_from
    return get_consolidated_report_rows(
        date_from=date_from,
        date_to=date_to,
        plating_stock_no=plating_stock_no,
        module=module,
    )


@login_required(login_url='login')
@require_admin
def consolidated_report_preview(request):
    """Paginated JSON preview — 10 rows per page, same data as download."""
    try:
        rows = _consolidated_rows_from_request(request)
    except Exception:
        logger.exception('Consolidated report preview failed')
        return JsonResponse({'error': 'Failed to build preview'}, status=500)

    paginator = Paginator(rows, 10)
    try:
        page_number = int(request.GET.get('page', 1))
    except (TypeError, ValueError):
        page_number = 1
    page = paginator.get_page(page_number)

    return JsonResponse({
        'results': list(page.object_list),
        'page': page.number,
        'num_pages': paginator.num_pages,
        'total_records': paginator.count,
        'has_previous': page.has_previous(),
        'has_next': page.has_next(),
    })


@login_required(login_url='login')
@require_admin
def consolidated_report_download(request):
    """Excel download for the consolidated journey report."""
    try:
        rows = _consolidated_rows_from_request(request)
    except Exception:
        logger.exception('Consolidated report download failed')
        return HttpResponse('Failed to build report', status=500)

    if not rows:
        return HttpResponse('No data found', status=404)

    excel_rows = [
        {
            'S.No': row['s_no'],
            'Plating Stock No.': row['plating_stk_no'],
            'Lot Qty': row['lot_qty'],
            **row['modules'],
            'Remarks': row['remarks'],
        }
        for row in rows
    ]
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        pd.DataFrame(excel_rows, columns=CONSOLIDATED_COLUMNS).to_excel(
            writer, sheet_name='Consolidated Report', index=False
        )
        _autofit_excel_columns(writer)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename=consolidated_report.xlsx'
    response['Content-Length'] = len(output.getvalue())
    output.close()
    return response


@login_required(login_url='login')
@require_admin
def plating_stock_autocomplete(request):
    """Partial-match autocomplete for Plating Stock No."""
    from .selectors import search_plating_stock

    try:
        results = search_plating_stock(request.GET.get('q', ''))
    except Exception:
        logger.exception('Plating stock autocomplete failed')
        results = []
    return JsonResponse({'results': results})
